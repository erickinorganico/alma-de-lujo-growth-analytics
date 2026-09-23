"""Private, aggregate-only client workbook review; no business executor."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_CEILING
from hashlib import sha256
from pathlib import Path
import csv
import json
import math
import re
import zipfile

ROOT = Path(__file__).resolve().parent.parent
CONTRACT = json.loads((ROOT / 'client/contract.json').read_text(encoding='utf-8'))
NUMERIC = {
    'CONFIG': ['observation_days', 'safety_days', 'target_margin', 'opening_cash', 'cash_floor', 'purchase_cap'],
    'CATALOGO': ['list_price', 'purchase_cost', 'inbound_freight', 'packaging', 'commission_rate', 'other_variable', 'lead_days', 'moq', 'pack_multiple'],
    'VENTAS': ['delivered_units', 'returned_units', 'restocked_units', 'net_revenue', 'variable_cost'],
    'STOCK': ['on_hand', 'reserved', 'in_transit', 'available_days', 'planned_qty'],
    'CAJA': ['amount'],
}
INTEGERS = {'observation_days', 'safety_days', 'lead_days', 'moq', 'pack_multiple', 'delivered_units', 'returned_units', 'restocked_units', 'on_hand', 'reserved', 'in_transit', 'available_days', 'planned_qty'}
DATE_FIELDS = {'as_of', 'date', 'count_date', 'eta', 'payment_date'}
MONEY_FIELDS = {'opening_cash', 'cash_floor', 'purchase_cap', 'list_price', 'purchase_cost', 'inbound_freight', 'packaging', 'other_variable', 'net_revenue', 'variable_cost', 'amount'}


def bounded(key, value):
    """Reject unsafe pasted magnitudes/precision before financial arithmetic."""
    if value is None:
        return True
    if key in DATE_FIELDS:
        return date(1900, 1, 1) <= value <= date(2100, 12, 31)
    if key in MONEY_FIELDS:
        return abs(value) <= 1_000_000_000 and value == value.quantize(Decimal('.01'))
    if key in ('lead_days', 'safety_days'):
        return abs(value) <= 3650
    if key in INTEGERS:
        return abs(value) <= 1_000_000
    return True


def number(value):
    if value is None or value == '' or isinstance(value, bool):
        return None
    try:
        d = Decimal(str(value))
        return d if d.is_finite() else None
    except (ValueError, ArithmeticError):
        return None


def day(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def clean_json(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dict):
        return {k: clean_json(v) for k, v in value.items()}
    if isinstance(value, list):
        return [clean_json(v) for v in value]
    return value


def read_inputs(path):
    """Read only declared inputs. Never trust cached workbook calculations."""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.stat().st_size > 20_000_000:
        raise ValueError('Se requiere un .xlsx de menos de 20 MB.')
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if sum(m.file_size for m in members) > 100_000_000 or len(members) > 1000:
            raise ValueError('El archivo supera los lÃ­mites de tamaÃ±o interno.')
        if any('vbaProject' in m.filename or m.filename.startswith('xl/externalLinks/') for m in members):
            raise ValueError('No se admiten macros ni vÃ­nculos a otros libros.')
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    required = {'CONFIG', *CONTRACT['sheets']}
    if not required <= set(wb.sheetnames):
        raise ValueError('Faltan hojas requeridas: ' + ', '.join(sorted(required - set(wb.sheetnames))))
    issues = []

    def issue(code, cell, message, severity='ERROR'):
        issues.append(dict(code=code, cell=cell, message=message, severity=severity))

    def read_cell(ws, coordinate):
        cell = ws[coordinate]
        if cell.data_type in {'f', 'e'}:
            issue('INPUT_FORMULA', ws.title + '!' + coordinate, 'Escribe un valor; esta celda de captura no admite fÃ³rmulas ni errores.')
            return None
        value = cell.value
        return value.strip() if isinstance(value, str) else value

    data = {'config': {}, 'catalog': [], 'sales': [], 'stock': [], 'cash': []}
    keymap = {'CATALOGO': 'catalog', 'VENTAS': 'sales', 'STOCK': 'stock', 'CAJA': 'cash'}

    def convert(sheet, key, value, cell):
        if value in (None, ''):
            return None
        if key in NUMERIC.get(sheet, []):
            n = number(value)
            if n is None or (key in INTEGERS and n != n.to_integral_value()):
                issue('INVALID_NUMBER', cell, 'Escribe un nÃºmero finito' + (' entero.' if key in INTEGERS else '.'))
                return None
            return n
        if key in DATE_FIELDS:
            d = day(value)
            if d is None:
                issue('INVALID_DATE', cell, 'Escribe una fecha vÃ¡lida, por ejemplo 2026-09-21.')
            return d
        if not isinstance(value, str) or len(value) > 120:
            issue('INVALID_TEXT', cell, 'Texto de hasta 120 caracteres.')
            return None
        if key == 'source_ref' and not re.fullmatch(r'[A-Za-z0-9._/-]{1,40}', value):
            issue('PRIVATE_REFERENCE', cell, 'Usa una referencia corta sin nombres, correo, telÃ©fono ni direcciÃ³n; por ejemplo CORTE-2026-09.')
        if key == 'sku' and not re.fullmatch(r'[A-Za-z0-9_-]{1,40}', value):
            issue('INVALID_SKU', cell, 'Usa letras, nÃºmeros, guion o guion bajo, sin espacios.')
        return value

    for key, cell in CONTRACT['parameters'].items():
        data['config'][key] = convert('CONFIG', key, read_cell(wb['CONFIG'], cell), 'CONFIG!' + cell)
    for sheet, columns in CONTRACT['sheets'].items():
        ws = wb[sheet]
        maximum = CONTRACT['last_sales_row'] if sheet == 'VENTAS' else CONTRACT['last_cash_row'] if sheet == 'CAJA' else CONTRACT['last_sku_row']
        if ws.max_row > 10000:
            raise ValueError('La hoja ' + sheet + ' tiene demasiadas filas.')
        for row in range(CONTRACT['first_row'], ws.max_row + 1):
            values = [read_cell(ws, f'{get_column_letter(col)}{row}') for col in range(1, len(columns) + 1)]
            if all(v in (None, '') for v in values):
                continue
            if row > maximum:
                issue('CAPACITY', f'{sheet}!A{row}', 'La fila excede la capacidad del kit; no se incorpora al anÃ¡lisis.')
                continue
            record = {'_row': row}
            for col, key in enumerate(columns, 1):
                record[key] = convert(sheet, key, values[col-1], f'{sheet}!{get_column_letter(col)}{row}')
            data[keymap[sheet]].append(record)
    data['metadata'] = {'contract_version': CONTRACT['version'], 'source_sha256': sha256(path.read_bytes()).hexdigest(), 'workbook_marker': wb['CONFIG']['B3'].value, 'private': True}
    wb.close()
    return data, issues


def evaluate(data, initial_issues=None):
    """Independent Decimal oracle for the customer workbook's declared model."""
    issues = list(initial_issues or [])
    config = data['config']
    cutoff = day(config.get('as_of'))
    config = {k: number(v) if k in NUMERIC['CONFIG'] else v for k, v in config.items()}
    window = config.get('observation_days')

    def issue(code, cell, message, severity='ERROR'):
        issues.append(dict(code=code, cell=cell, message=message, severity=severity))

    for key in ['observation_days', 'safety_days', 'target_margin', 'opening_cash', 'cash_floor', 'purchase_cap']:
        value = config.get(key)
        if value is None:
            issue('MISSING_PARAMETER', 'CONFIG!' + CONTRACT['parameters'][key], 'Completa el parÃ¡metro para habilitar los cÃ¡lculos que dependen de Ã©l.', 'WARNING')
        elif (key in INTEGERS and value != value.to_integral_value()) or (key != 'opening_cash' and value < 0) or (key == 'observation_days' and not 1 <= value <= 365) or (key == 'target_margin' and value >= 1):
            issue('INVALID_PARAMETER', 'CONFIG!' + CONTRACT['parameters'][key], 'El parÃ¡metro estÃ¡ fuera del rango permitido.')
    if cutoff is None:
        issue('MISSING_CUTOFF', 'CONFIG!B5', 'Completa la fecha de corte.', 'WARNING')
    elif not bounded('as_of', cutoff):
        issue('DATE_RANGE', 'CONFIG!B5', 'Usa una fecha entre 1900 y 2100.')
    for key, value in config.items():
        if key in NUMERIC['CONFIG'] and not bounded(key, value):
            issue('INPUT_RANGE', 'CONFIG!' + CONTRACT['parameters'][key], 'Revisa el rango y la precisiÃ³n: importes en pesos con mÃ¡ximo dos decimales; plazos de hasta 3650 dÃ­as.')
    for key in ['policy_reviewed', 'cash_plan_complete', 'sales_complete']:
        if config.get(key) not in ('SI', 'NO', None):
            issue('INVALID_CONFIRMATION', 'CONFIG!' + CONTRACT['parameters'][key], 'Selecciona SI o NO.')
    for table, sheet in [('catalog','CATALOGO'),('stock','STOCK'),('sales','VENTAS'),('cash','CAJA')]:
        for record in data[table]:
            for key in NUMERIC[sheet]:
                record[key] = number(record.get(key))
            for key in DATE_FIELDS & set(record):
                record[key] = day(record[key])
            for key in set(NUMERIC[sheet]) | (DATE_FIELDS & set(record)):
                if not bounded(key, record.get(key)):
                    issue('INPUT_RANGE', f"{sheet}!A{record.get('_row','?')}", 'Revisa el rango de ' + key + ': importes de hasta mil millones con dos decimales, cantidades de hasta un millÃ³n, plazos de hasta 3650 dÃ­as y fechas de 1900 a 2100.')

    if any(x['severity']=='ERROR' for x in issues):
        return clean_json(dict(version='3.0',status='CORREGIR_CAPTURA',issues=issues,skus=[],cash_weeks=[],external_execution='PROHIBITED',metadata=data.get('metadata',{})))

    catalog = {}; stocks = {}; sale_keys = set()
    for row in data['catalog']:
        ref = f"CATALOGO!A{row.get('_row', '?')}"; sku = row.get('sku')
        if not sku or sku in catalog:
            issue('DUPLICATE_SKU', ref, 'Cada SKU debe ser Ãºnico y no estar vacÃ­o.')
        catalog[sku] = row
        if row.get('status') not in CONTRACT['enums']['product_status']:
            issue('PRODUCT_STATUS', ref, 'Estado de producto invÃ¡lido.')
        for key in NUMERIC['CATALOGO']:
            n = row.get(key)
            if n is not None and (n < 0 or (key in ('moq','pack_multiple','list_price') and n <= 0) or (key == 'commission_rate' and n >= 1) or (key in INTEGERS and n != n.to_integral_value())):
                issue('CATALOG_RANGE', ref, 'Revisa el valor de ' + key + '.')
    for row in data['stock']:
        ref = f"STOCK!A{row.get('_row', '?')}"; sku = row.get('sku')
        if sku not in catalog or sku in stocks:
            issue('STOCK_SKU', ref, 'El SKU debe existir en catÃ¡logo y tener un Ãºnico corte.')
        stocks[sku] = row
        for key in NUMERIC['STOCK']:
            n = row.get(key)
            if n is not None and (n < 0 or n != n.to_integral_value()):
                issue('STOCK_RANGE', ref, 'Las cantidades y dÃ­as deben ser enteros no negativos.')
        if row.get('reserved') is not None and row.get('on_hand') is not None and row['reserved'] > row['on_hand']:
            issue('RESERVED_EXCEEDS_STOCK', ref, 'La reserva supera la existencia fÃ­sica.')
        if row.get('available_days') is not None and window is not None and row['available_days'] > window:
            issue('EXPOSURE_RANGE', ref, 'Los dÃ­as disponibles superan la ventana de observaciÃ³n.')
        if row.get('transit_confirmed') not in ('SI','NO',None):
            issue('TRANSIT_CONFIRMATION', ref, 'Selecciona SI o NO para el trÃ¡nsito.')
    for row in data['sales']:
        ref = f"VENTAS!A{row.get('_row', '?')}"; key = (row.get('date'),row.get('sku'))
        if key in sale_keys:
            issue('DUPLICATE_SALES_AGGREGATE', ref, 'Consolida una sola fila por fecha y SKU; no dupliques ventas.')
        sale_keys.add(key)
        if row.get('sku') not in catalog or row.get('date') is None or (cutoff and row['date'] > cutoff):
            issue('SALES_REFERENCE', ref, 'Revisa SKU y fecha; no incluyas ventas futuras.')
        for k in ['delivered_units','returned_units','restocked_units','net_revenue']:
            n = row.get(k)
            if n is None or (k != 'net_revenue' and (n < 0 or n != n.to_integral_value())):
                issue('SALES_VALUE', ref, 'Completa cantidades vÃ¡lidas e ingreso neto; cero debe ser explÃ­cito.')
        if row.get('restocked_units') is not None and row.get('returned_units') is not None and row['restocked_units'] > row['returned_units']:
            issue('RESTOCK_EXCEEDS_RETURN', ref, 'Las unidades reintegradas no pueden superar las devueltas.')
    for row in data['cash']:
        ref = f"CAJA!A{row.get('_row', '?')}"
        if row.get('date') is None or row.get('direction') not in ('ENTRADA','SALIDA') or row.get('amount') is None or row['amount'] <= 0:
            issue('CASH_VALUE', ref, 'Completa fecha, direcciÃ³n y monto positivo.')
        elif cutoff and not cutoff < row['date'] <= cutoff + timedelta(days=91):
            issue('CASH_DATE', ref, 'El plan de caja debe cubrir desde el dÃ­a posterior al corte hasta el dÃ­a 91.')
    if any(x['severity']=='ERROR' for x in issues):
        return clean_json(dict(version='3.0',status='CORREGIR_CAPTURA',issues=issues,skus=[],cash_weeks=[],external_execution='PROHIBITED',metadata=data.get('metadata',{})))
    if not catalog:
        issue('EMPTY_CATALOG','CATALOGO!A6','Agrega tus SKUs para comenzar.','WARNING')
    start = cutoff - timedelta(days=int(window)-1) if cutoff and window and window > 0 else None
    selected = [r for r in data['sales'] if start and start <= r['date'] <= cutoff]
    rows=[]; purchase_events=[]; planned_total=Decimal(0)
    purchase_fields=['list_price','purchase_cost','inbound_freight','packaging','commission_rate','other_variable','lead_days','moq','pack_multiple']
    plan_known=bool(catalog) and all(config.get(k) is not None for k in ['observation_days','safety_days','target_margin']) and all(all(item.get(k) is not None for k in purchase_fields) for item in catalog.values())
    for sku, item in catalog.items():
        stock=stocks.get(sku,{})
        sale_rows=[r for r in selected if r['sku']==sku]
        observed = config.get('sales_complete')=='SI' and start is not None
        units=sum((r['delivered_units'] for r in sale_rows),Decimal(0)) if observed else None
        restock=sum((r['restocked_units'] for r in sale_rows),Decimal(0)) if observed else None
        revenue=sum((r['net_revenue'] for r in sale_rows),Decimal(0)) if observed else None
        cost_known=observed and all(r['variable_cost'] is not None for r in sale_rows)
        actual_variable=sum((r['variable_cost'] for r in sale_rows),Decimal(0)) if cost_known else None
        cost_keys=['purchase_cost','inbound_freight','packaging','other_variable']
        unit_cost=sum((item[k] for k in cost_keys),Decimal(0)) if all(item.get(k) is not None for k in cost_keys) else None
        commission=item.get('commission_rate'); target=config.get('target_margin'); price=item.get('list_price')
        denominator=1-commission-target if commission is not None and target is not None else None
        floor_price=(unit_cost/denominator).quantize(Decimal('.01'),rounding=ROUND_CEILING) if unit_cost is not None and denominator is not None and denominator>0 else None
        unit_contribution=price*(1-commission)-unit_cost if price and commission is not None and unit_cost is not None else None
        margin=unit_contribution/price if unit_contribution is not None else None
        available=stock['on_hand']-stock['reserved'] if stock.get('on_hand') is not None and stock.get('reserved') is not None else None
        reasons=[]; lead=item.get('lead_days'); safety=config.get('safety_days'); transit=stock.get('in_transit'); eta=stock.get('eta'); eligible=Decimal(0)
        if item.get('status')!='ACTIVO':reasons.append('PRODUCTO_NO_ACTIVO')
        if config.get('policy_reviewed')!='SI':reasons.append('POLITICA_POR_REVISAR')
        if not observed:reasons.append('VENTAS_INCOMPLETAS')
        if not cutoff or stock.get('count_date')!=cutoff:reasons.append('ACTUALIZAR_CONTEO')
        if stock.get('available_days')!=window or window is None:reasons.append('DISPONIBILIDAD_NO_OBSERVADA_COMPLETA')
        if any(v is None for v in [available,lead,safety,transit,item.get('moq'),item.get('pack_multiple'),unit_cost,commission,price,target]):reasons.append('DATOS_FALTANTES')
        if denominator is not None and denominator<=0:reasons.append('MARGEN_OBJETIVO_INVIABLE')
        if margin is not None and target is not None and margin<target:reasons.append('REVISAR_PRECIO_O_COSTO')
        if transit and transit>0:
            if eta is None or not cutoff or eta<=cutoff:reasons.append('REVISAR_TRANSITO')
            elif lead is not None and stock.get('transit_confirmed')=='SI' and eta<=cutoff+timedelta(days=int(lead)):
                eligible=transit
            elif stock.get('transit_confirmed')!='SI':reasons.append('TRANSITO_NO_CONFIRMADO')
            else:reasons.append('TRANSITO_FUERA_DEL_PLAZO')
        rate=max(Decimal(0),units-restock)/window if observed and window else None
        cover=available/rate if available is not None and rate else None
        suggestion=None
        if not reasons:
            needed=max(0,math.ceil(rate*(lead+safety)-available-eligible))
            suggestion=math.ceil(max(needed,item['moq'])/item['pack_multiple'])*int(item['pack_multiple']) if needed else 0
        chosen=stock.get('planned_qty'); pay_date=stock.get('payment_date'); landed=item.get('purchase_cost')+item.get('inbound_freight') if item.get('purchase_cost') is not None and item.get('inbound_freight') is not None else None
        plan_cost=None
        if chosen is None:
            plan_known=False
        elif chosen>0:
            if landed is None or item.get('moq') is None or item.get('pack_multiple') is None or chosen<item['moq'] or chosen%item['pack_multiple']!=0 or not cutoff or pay_date is None or not cutoff<pay_date<=cutoff+timedelta(days=91):
                plan_known=False;reasons.append('CORREGIR_PLAN_DE_COMPRA')
            else:
                plan_cost=chosen*landed;planned_total+=plan_cost;purchase_events.append((pay_date,plan_cost))
        else:plan_cost=Decimal(0)
        rows.append(dict(sku=sku,product=item.get('product'),color=item.get('color'),size=item.get('size'),status='REVISAR' if reasons else 'PROPUESTA_CONDICIONAL' if suggestion else 'SIN_REPOSICION',reasons=reasons,delivered_units=units,restocked_units=restock,net_revenue_mxn=revenue,variable_cost_mxn=actual_variable,contribution_mxn=revenue-actual_variable if revenue is not None and actual_variable is not None else None,unit_variable_cost_mxn=unit_cost,unit_contribution_mxn=unit_contribution,unit_margin=margin,target_price_mxn=floor_price,available_units=available,eligible_transit_units=eligible,daily_consumption=rate,cover_days=cover,suggested_qty=suggestion,planned_qty=chosen,planned_purchase_mxn=plan_cost))
    cash_known=bool(cutoff and config.get('cash_plan_complete')=='SI' and all(config.get(k) is not None for k in ['opening_cash','cash_floor','purchase_cap']))
    weeks=[];days=[];base=config.get('opening_cash');after=base;minbase=base;minafter=after
    if cash_known:
        for i in range(91):
            current=cutoff+timedelta(days=1+i)
            incoming=sum((r['amount'] for r in data['cash'] if current==r['date'] and r['direction']=='ENTRADA'),Decimal(0))
            outgoing=sum((r['amount'] for r in data['cash'] if current==r['date'] and r['direction']=='SALIDA'),Decimal(0))
            purchases=sum((amount for when,amount in purchase_events if current==when),Decimal(0)) if plan_known else None
            base+=incoming-outgoing;minbase=min(minbase,base)
            if plan_known:after+=incoming-outgoing-purchases;minafter=min(minafter,after)
            days.append(dict(date=current,in_mxn=incoming,out_existing_mxn=outgoing,planned_purchases_mxn=purchases,base_closing_mxn=base,planned_closing_mxn=after if plan_known else None))
        for i in range(13):
            group=days[i*7:i*7+7]
            minimum=min(r['planned_closing_mxn'] for r in group) if plan_known else None
            weeks.append(dict(week=i+1,start=group[0]['date'],end=group[-1]['date'],in_mxn=sum(r['in_mxn'] for r in group),out_existing_mxn=sum(r['out_existing_mxn'] for r in group),planned_purchases_mxn=sum(r['planned_purchases_mxn'] for r in group) if plan_known else None,base_closing_mxn=group[-1]['base_closing_mxn'],planned_closing_mxn=group[-1]['planned_closing_mxn'],min_daily_base_mxn=min(r['base_closing_mxn'] for r in group),min_daily_after_mxn=minimum,below_floor=(minimum<config['cash_floor']) if plan_known else None))
    budget=max(Decimal(0),min(config['purchase_cap'],minbase-config['cash_floor'])) if cash_known else None
    budget_status='DATOS_PENDIENTES' if not cash_known or not plan_known else 'EXCEDE_LIMITE' if planned_total>config['purchase_cap'] or minafter<config['cash_floor'] else 'DENTRO_DEL_ESCENARIO'
    return clean_json(dict(version='3.0',status='DATOS_PENDIENTES' if not catalog or not start else 'REVISAR',as_of=cutoff,issues=issues,skus=rows,cash_weeks=weeks,cash_days=days,budget=dict(status=budget_status,available_before_proposals_mxn=budget,planned_purchase_mxn=planned_total if plan_known else None,min_daily_cash_after_mxn=minafter if cash_known and plan_known else None,cash_floor_mxn=config.get('cash_floor')),external_execution='PROHIBITED',metadata=data.get('metadata',{}),limits=['Propuestas condicionales; no son Ã³rdenes ni autorizaciones.','Caja por cierres diarios: puede haber faltantes dentro del dÃ­a.','Costos histÃ³ricos opcionales no se sustituyen por costos actuales del catÃ¡logo.','La demanda observada y el precio objetivo no prueban preferencia ni elasticidad.']))


def review_file(input_path, output):
    """Write only to ignored .local/client-runs; never publish client workbooks."""
    output=Path(output).resolve();allowed=(ROOT/'.local/client-runs').resolve()
    if not output.is_relative_to(allowed) or output==allowed:
        raise ValueError('Guarda el informe en una subcarpeta de .local/client-runs para mantenerlo fuera del repositorio pÃºblico.')
    if output.exists() and any(output.iterdir()):
        raise ValueError('La carpeta ya contiene datos. Usa un nombre de corte nuevo.')
    data,issues=read_inputs(input_path);report=evaluate(data,issues)
    output.mkdir(parents=True,exist_ok=True)
    content=json.dumps(report,ensure_ascii=False,indent=2).encode('utf-8');(output/'informe.json').write_bytes(content)
    digest=sha256(content).hexdigest()
    from alma.client_report import render
    (output/'informe.html').write_text(render(report, digest), encoding='utf-8')
    with (output/'correcciones.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=['severity','code','cell','message']);writer.writeheader();writer.writerows(report['issues'])
    (output/'PARA_EL_ANALISTA.md').write_text(f'''# RevisiÃ³n nativa del corte\n\nLee informe.json en esta misma carpeta. Su SHA-256 es `{digest}`.\nVerifica este hash antes de analizarlo. Usa sÃ³lo este corte y referencias JSON pointer exactas.\nSepara observaciones, incertidumbres e hipÃ³tesis; no conviertas campos null en cero.\nRevisa como mÃ¡ximo tres decisiones de inventario/precio/caja y explica alternativa, mÃ©trica,\nventana, poblaciÃ³n y condiciÃ³n de cierre. La autoridad de ejecuciÃ³n externa es PROHIBITED.\nUn agente distinto debe revisar la respuesta. No subas el libro ni este informe a GitHub.\nNo completes datos faltantes con el ejemplo sintÃ©tico.\n\nEl generador preparÃ³ esta solicitud; no realizÃ³ una llamada a un modelo.\n''',encoding='utf-8')
    return report
