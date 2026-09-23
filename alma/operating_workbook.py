"""Fail-closed XLSX adapter for the aggregate ``operating-v1`` contract."""
from __future__ import annotations

import csv
import hashlib
import json
import math
import os
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any

from openpyxl import load_workbook

from .operating_contracts import (
    CONTRACT_VERSION, SOURCE_NAMES, SOURCES, OperatingContractError, canonical_json,
    columns_for, validate_metadata,
)
from .operating_interchange import MAX_CSV_BYTES, MAX_ROWS_PER_SOURCE, parse_pack


MAX_WORKBOOK_BYTES = 20_000_000
MAX_ARCHIVE_BYTES = 100_000_000
MAX_ARCHIVE_MEMBERS = 2_000
MAX_INPUT_ROWS = 1_000
SUPPORT_SHEETS = ("INICIO", "DICCIONARIO", "COMPLETITUD")
_EXTERNAL_REL = re.compile(rb'TargetMode=["\']External["\']', re.IGNORECASE)
_PII_LIKE = re.compile(r"(?:[^@\s]+@[^@\s]+\.[^@\s]+|\b\d{7,15}\b)")


class WorkbookContractError(ValueError):
    """A sanitized workbook diagnostic with an actionable correction."""

    def __init__(self, code: str, *, sheet: str = "workbook", row: int | None = None,
                 field: str | None = None, issue: str | None = None,
                 correction: str = "Corrige el libro y exporta a una carpeta nueva.",
                 upstream_code: str | None = None) -> None:
        self.code = code
        self.sheet = sheet
        self.row = row
        self.field = field
        self.issue = issue or code
        self.correction = correction
        self.upstream_code = upstream_code
        location = sheet + (f":{row}" if row is not None else "") + (f":{field}" if field else "")
        super().__init__(f"{code} at {location}: {self.issue}. {correction}")

    def as_dict(self) -> dict[str, Any]:
        return {
            "code": self.code,
            "sheet": self.sheet,
            "row": self.row,
            "field": self.field,
            "issue": self.issue,
            "correction": self.correction,
            "upstream_code": self.upstream_code,
        }


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fail(code: str, *, sheet: str = "workbook", row: int | None = None,
          field: str | None = None, issue: str | None = None,
          correction: str = "Corrige el libro y exporta a una carpeta nueva.",
          upstream_code: str | None = None) -> None:
    raise WorkbookContractError(code, sheet=sheet, row=row, field=field, issue=issue,
                                correction=correction, upstream_code=upstream_code)


def _inside(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True


def _inspect_archive(path: Path) -> None:
    if path.suffix.lower() != ".xlsx" or path.is_symlink():
        _fail("workbook.type", correction="Usa un archivo .xlsx local, sin enlaces simbólicos.")
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise WorkbookContractError("workbook.missing", correction="Selecciona un libro existente.") from exc
    if size > MAX_WORKBOOK_BYTES:
        _fail("workbook.size", correction="Reduce el libro a menos de 20 MB.")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_MEMBERS or sum(item.file_size for item in members) > MAX_ARCHIVE_BYTES:
                _fail("workbook.archive_size", correction="Elimina contenido incrustado y usa la plantilla oficial.")
            seen: set[str] = set()
            for member in members:
                name = member.filename.replace("\\", "/")
                pure = PurePosixPath(name)
                if (name in seen or pure.is_absolute() or ".." in pure.parts or member.flag_bits & 1
                        or (member.external_attr >> 16) & 0o170000 == 0o120000):
                    _fail("workbook.archive_member", correction="Regenera el archivo desde la plantilla oficial.")
                seen.add(name)
                lowered = name.lower()
                if ("vbaproject" in lowered or lowered.startswith("xl/externallinks/")
                        or lowered.startswith("xl/embeddings/") or lowered.startswith("customxml/")
                        or lowered.startswith("xl/connections")):
                    _fail("workbook.external_content", correction="Elimina macros, vínculos y objetos externos.")
                if lowered.endswith(".rels") and _EXTERNAL_REL.search(archive.read(member)):
                    _fail("workbook.external_content", correction="Elimina relaciones a recursos externos.")
    except zipfile.BadZipFile as exc:
        raise WorkbookContractError("workbook.archive", correction="Usa un .xlsx válido.") from exc


def _metadata(workbook: Any) -> dict[str, Any]:
    start = workbook["INICIO"]
    if start["A3"].value != "Contrato" or start["A4"].value != "Clase" or start["A5"].value != "Corte" \
            or start["A6"].value != "Zona horaria" or start["A7"].value != "Moneda":
        _fail("workbook.metadata_layout", sheet="INICIO", correction="Restaura las etiquetas de la plantilla oficial.")
    completeness = workbook["COMPLETITUD"]
    expected = ["relacion", "dominio", "clasificacion", "estado_si_falta", "familias_metricas",
                "coverage_status", "window_start", "window_end", "regla"]
    actual = [completeness.cell(1, column).value for column in range(1, len(expected) + 1)]
    if actual != expected:
        _fail("workbook.coverage_header", sheet="COMPLETITUD", row=1,
              correction="Restaura las columnas de cobertura de la plantilla oficial.")
    coverage: dict[str, dict[str, Any]] = {}
    for index, relation in enumerate(SOURCE_NAMES, start=2):
        if completeness.cell(index, 1).value != relation:
            _fail("workbook.coverage_relation", sheet="COMPLETITUD", row=index, field="relacion",
                  correction="Conserva las 22 relaciones en el orden original.")
        coverage[relation] = {
            "status": completeness.cell(index, 6).value,
            "window_start": completeness.cell(index, 7).value,
            "window_end": completeness.cell(index, 8).value,
        }
    if any(completeness.cell(row, 1).value not in (None, "") for row in range(24, completeness.max_row + 1)):
        _fail("workbook.coverage_relation", sheet="COMPLETITUD", row=24,
              correction="Elimina relaciones adicionales de la matriz.")
    metadata = {
        "contract_version": start["B3"].value,
        "input_class": start["B4"].value,
        "cutoff_at": start["B5"].value,
        "timezone": start["B6"].value,
        "currency": start["B7"].value,
        "coverage": coverage,
    }
    try:
        validate_metadata(metadata)
    except OperatingContractError as exc:
        _fail("workbook.metadata", sheet="INICIO", issue=exc.code,
              correction="Completa clase, corte, zona, moneda y cobertura según el contrato.",
              upstream_code=exc.code)
    return metadata


def _inspect_surfaces(workbook: Any) -> None:
    expected = [*SUPPORT_SHEETS, *SOURCE_NAMES]
    if workbook.sheetnames != expected:
        _fail("workbook.sheet_set", correction="Conserva sólo las hojas oficiales y en su orden original.")
    for sheet in workbook.worksheets:
        if sheet.sheet_state != "visible":
            _fail("workbook.hidden_sheet", sheet=sheet.title, correction="Haz visible la hoja o restaura la plantilla.")
        for row_number, dimension in sheet.row_dimensions.items():
            if dimension.hidden:
                _fail("workbook.hidden_row", sheet=sheet.title, row=row_number,
                      correction="Haz visible la fila antes de exportar.")
        for column, dimension in sheet.column_dimensions.items():
            if dimension.hidden:
                _fail("workbook.hidden_column", sheet=sheet.title, field=column,
                      correction="Haz visible la columna antes de exportar.")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    _fail("workbook.comment", sheet=sheet.title, row=cell.row, field=cell.column_letter,
                          correction="Mueve notas a evidencia autorizada y elimina comentarios.")
                if sheet.title in SOURCE_NAMES and cell.hyperlink is not None:
                    _fail("workbook.hyperlink", sheet=sheet.title, row=cell.row, field=cell.column_letter,
                          correction="Elimina el vínculo de la celda de captura.")
    for name in workbook.defined_names.values():
        if not name.name.startswith("_xlnm."):
            _fail("workbook.defined_name", correction="Elimina nombres personalizados del libro.")
    properties = workbook.properties
    for value in (properties.creator, properties.lastModifiedBy, properties.keywords,
                  properties.description, properties.identifier, properties.category):
        if isinstance(value, str) and _PII_LIKE.search(value):
            _fail("workbook.metadata_pii", correction="Elimina correo, teléfono o identificadores personales de las propiedades.")


def _cell_text(value: Any, field: dict[str, Any], relation: str, row: int) -> str:
    name = field["name"]
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        _fail("workbook.value", sheet=relation, row=row, field=name, issue="boolean not accepted",
              correction="Escribe el valor exacto declarado por el diccionario.")
    kind = field["type"]
    if kind in {"integer", "nonnegative_integer", "signed_integer"}:
        try:
            decimal = Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise WorkbookContractError("workbook.value", sheet=relation, row=row, field=name,
                                        issue="integer required", correction="Escribe un entero exacto.",
                                        upstream_code="value.integer") from exc
        if not decimal.is_finite() or decimal != decimal.to_integral_value():
            _fail("workbook.value", sheet=relation, row=row, field=name, issue="integer required",
                  correction="Escribe un entero exacto; los importes se capturan en centavos.",
                  upstream_code="value.integer")
        return format(decimal.quantize(Decimal(1)), "f")
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if not isinstance(value, str):
            _fail("workbook.value", sheet=relation, row=row, field=name, issue="date required",
                  correction="Usa una fecha ISO AAAA-MM-DD.", upstream_code="value.date")
        return value
    if kind == "timestamp":
        if isinstance(value, datetime):
            if value.tzinfo is None:
                _fail("workbook.value", sheet=relation, row=row, field=name, issue="timezone required",
                      correction="Usa una marca ISO con desplazamiento horario.", upstream_code="value.timestamp")
            return value.isoformat()
        if not isinstance(value, str):
            _fail("workbook.value", sheet=relation, row=row, field=name, issue="timestamp required",
                  correction="Usa una marca ISO con desplazamiento horario.", upstream_code="value.timestamp")
        return value
    if not isinstance(value, str):
        _fail("workbook.value", sheet=relation, row=row, field=name, issue="text required",
              correction="Escribe texto conforme al diccionario.", upstream_code="value.token")
    return value


def _read_rows(workbook: Any) -> dict[str, list[dict[str, str]]]:
    tables: dict[str, list[dict[str, str]]] = {}
    for relation in SOURCE_NAMES:
        sheet = workbook[relation]
        fields = SOURCES[relation]["fields"]
        expected = list(columns_for(relation))
        actual = [sheet.cell(6, column).value for column in range(1, len(expected) + 1)]
        if actual != expected or len(set(actual)) != len(actual):
            _fail("workbook.header", sheet=relation, row=6,
                  correction="Restaura exactamente los encabezados y su orden.")
        if any(sheet.cell(6, column).value not in (None, "") for column in range(len(expected) + 1, sheet.max_column + 1)):
            _fail("workbook.header", sheet=relation, row=6,
                  correction="Elimina columnas adicionales.")
        if sheet.max_row > 6 + MAX_INPUT_ROWS:
            _fail("workbook.rows", sheet=relation, row=sheet.max_row,
                  correction=f"Usa como máximo {MAX_INPUT_ROWS} filas por relación.")
        rows: list[dict[str, str]] = []
        for row_number in range(7, sheet.max_row + 1):
            cells = [sheet.cell(row_number, column) for column in range(1, len(expected) + 1)]
            if any(cell.data_type in {"f", "e"} for cell in cells):
                cell = next(cell for cell in cells if cell.data_type in {"f", "e"})
                _fail("workbook.formula", sheet=relation, row=row_number, field=cell.column_letter,
                      correction="Reemplaza la fórmula o error por un valor de captura.")
            if all(cell.value in (None, "") for cell in cells):
                continue
            rows.append({field["name"]: _cell_text(cell.value, field, relation, row_number)
                         for field, cell in zip(fields, cells, strict=True)})
        tables[relation] = rows
    return tables


def _write_pack(stage: Path, metadata: dict[str, Any], tables: dict[str, list[dict[str, str]]]) -> Path:
    pack = stage / "source-pack"
    pack.mkdir()
    (pack / "metadata.json").write_bytes(canonical_json(metadata) + b"\n")
    for relation in SOURCE_NAMES:
        path = pack / f"{relation}.csv"
        with path.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns_for(relation), lineterminator="\n",
                                    extrasaction="raise")
            writer.writeheader()
            writer.writerows(tables[relation])
        if path.stat().st_size > MAX_CSV_BYTES or len(tables[relation]) > MAX_ROWS_PER_SOURCE:
            _fail("workbook.export_limit", sheet=relation,
                  correction="Reduce la captura dentro de los límites del contrato.")
    return pack


def _upstream_error(error: OperatingContractError) -> WorkbookContractError:
    parts = error.location.split(":")
    sheet = parts[0].removesuffix(".csv") if parts else "workbook"
    row = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else None
    field = parts[2] if len(parts) > 2 else None
    return WorkbookContractError(
        "workbook.intake_rejected", sheet=sheet, row=row, field=field,
        issue=error.code, upstream_code=error.code,
        correction="Corrige la fila indicada según el contrato y vuelve a exportar.",
    )


def _export_manifest(workbook_hash: str, pack: Path, parsed: dict[str, Any]) -> dict[str, Any]:
    relations: dict[str, Any] = {}
    for relation in SOURCE_NAMES:
        contract = SOURCES[relation]
        typed_rows = parsed["tables"][relation]
        numeric_fields = {field["name"] for field in contract["fields"]
                          if field["type"] in {"integer", "nonnegative_integer", "signed_integer"}}
        relations[relation] = {
            "filename": contract["filename"],
            "row_count": len(typed_rows),
            "null_count": sum(value is None for row in typed_rows for value in row.values()),
            "zero_count": sum(row[name] == 0 for row in typed_rows for name in numeric_fields),
            "cents_fields": [field["name"] for field in contract["fields"] if field["unit"] == "MXN cents"],
            "units": {field["name"]: field["unit"] for field in contract["fields"]},
            "sha256": parsed["source_sha256"][contract["filename"]],
        }
    return {
        "manifest_version": "workbook-source-pack-export-v1",
        "status": "PASS",
        "contract_version": CONTRACT_VERSION,
        "workbook_sha256": workbook_hash,
        "metadata_sha256": parsed["metadata_sha256"],
        "normalized_rows_digest": parsed["normalized_rows_digest"],
        "relations": relations,
        "relation_order": list(SOURCE_NAMES),
        "preserves_unknown_zero": True,
        "money_boundary": "integer MXN cents",
    }


def export_workbook_to_pack(workbook_path: str | Path, output: str | Path, *,
                            private_root: str | Path) -> dict[str, Any]:
    """Export one XLSX into a new validated canonical source-pack bundle."""

    source = Path(workbook_path).resolve(strict=True)
    root = Path(private_root).resolve(strict=True)
    destination = Path(output).resolve(strict=False)
    if not _inside(destination, root) or destination == root:
        _fail("output.private_root", correction="Exporta dentro de una subcarpeta nueva del directorio privado autorizado.")
    if destination.exists():
        _fail("output.exists", correction="Usa una carpeta de exportación nueva; no se sobrescriben cortes.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    original_hash = _hash(source)
    _inspect_archive(source)
    workbook = load_workbook(source, data_only=False, read_only=False, keep_links=False)
    stage = Path(tempfile.mkdtemp(prefix=".workbook-export-", dir=destination.parent))
    try:
        _inspect_surfaces(workbook)
        metadata = _metadata(workbook)
        if metadata["input_class"] == "PRIVATE" and not _inside(source, root):
            _fail("workbook.private_root", correction="Guarda el libro privado dentro del directorio privado autorizado.")
        tables = _read_rows(workbook)
        pack = _write_pack(stage, metadata, tables)
        try:
            parsed = parse_pack(pack, private_root=root if metadata["input_class"] == "PRIVATE" else None)
        except OperatingContractError as exc:
            raise _upstream_error(exc) from exc
        manifest = _export_manifest(original_hash, pack, parsed)
        manifest_path = stage / "export-manifest.json"
        manifest_path.write_bytes(canonical_json(manifest) + b"\n")
        if _hash(source) != original_hash:
            _fail("workbook.mutated", correction="Restaura el libro original; el adaptador nunca debe modificarlo.")
        os.replace(stage, destination)
        return {
            "status": "PASS",
            "contract_version": CONTRACT_VERSION,
            "workbook_sha256": original_hash,
            "relations": list(SOURCE_NAMES),
            "bundle_path": str(destination),
            "pack_path": str(destination / "source-pack"),
            "manifest_path": str(destination / "export-manifest.json"),
            "manifest_sha256": _hash(destination / "export-manifest.json"),
        }
    except WorkbookContractError:
        raise
    finally:
        workbook.close()
        if stage.exists():
            shutil.rmtree(stage)


__all__ = ["WorkbookContractError", "export_workbook_to_pack"]
