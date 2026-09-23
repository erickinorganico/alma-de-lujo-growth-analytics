"""Independent, hand-calculated customer-kit semantics and hostile-input checks."""
from __future__ import annotations

from copy import deepcopy
from datetime import date, timedelta
from hashlib import sha256
from pathlib import Path
import sys
import importlib.util
import tempfile
import unittest
from unittest.mock import patch
import zipfile

sys.path.insert(0, str(Path(__file__).parents[1]))
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from alma import client_review as cr


def scenario():
    """Sunday cutoff; 14 deliveries over 7 fully observed days = 2/day."""
    cutoff = date(2026, 9, 20)
    return {
        'config': dict(as_of=cutoff, observation_days=7, safety_days=2,
                       target_margin=.4, opening_cash=1000, cash_floor=200,
                       purchase_cap=800, policy_reviewed='SI',
                       cash_plan_complete='SI', sales_complete='SI'),
        'catalog': [dict(_row=6, sku='SYN-SOCK-M', product='Synthetic sock',
                         category='Pilates', color='Rose', size='M', status='ACTIVO',
                         list_price=100, purchase_cost=20, inbound_freight=2,
                         packaging=3, commission_rate=.1, other_variable=0,
                         lead_days=5, moq=10, pack_multiple=6)],
        'sales': [dict(_row=6, date=cutoff, sku='SYN-SOCK-M', delivered_units=14,
                       returned_units=0, restocked_units=0, net_revenue=1400,
                       variable_cost=490, source_ref='SYN-001')],
        'stock': [dict(_row=6, sku='SYN-SOCK-M', count_date=cutoff,
                       on_hand=6, reserved=2, in_transit=0, eta=None,
                       transit_confirmed='NO', available_days=7,
                       planned_qty=0, payment_date=None)],
        'cash': [],
        'metadata': {'synthetic': True},
    }


def write_inputs(path, data):
    wb = Workbook()
    wb.active.title = 'CONFIG'
    wb['CONFIG']['B3'] = 'SYNTHETIC TEST FIXTURE'
    for key, cell in cr.CONTRACT['parameters'].items():
        wb['CONFIG'][cell] = data['config'].get(key)
    table_keys = dict(CATALOGO='catalog', VENTAS='sales', STOCK='stock', CAJA='cash')
    for name, columns in cr.CONTRACT['sheets'].items():
        ws = wb.create_sheet(name)
        for record in data[table_keys[name]]:
            for col, key in enumerate(columns, 1):
                ws.cell(record['_row'], col, record.get(key))
    wb.save(path)
    wb.close()


class ClientArithmeticTests(unittest.TestCase):
    def test_hand_calculated_quantity_price_margin_and_purchase(self):
        data = scenario()
        data['stock'][0].update(planned_qty=12, payment_date=date(2026, 9, 21))
        result = cr.evaluate(data)
        row = result['skus'][0]
        # (14/7)*(5+2) - (6-2) = 10; MOQ 10, pack 6 => 12.
        self.assertEqual(row['suggested_qty'], 12)
        self.assertEqual(row['available_units'], 4)
        self.assertEqual(row['daily_consumption'], 2)
        self.assertEqual(row['cover_days'], 2)
        # Unit costs 20+2+3+0=25; contribution 100*.9-25=65.
        self.assertEqual(row['unit_variable_cost_mxn'], 25)
        self.assertEqual(row['unit_contribution_mxn'], 65)
        self.assertEqual(row['unit_margin'], .65)
        self.assertEqual(row['target_price_mxn'], 50)
        self.assertEqual(row['contribution_mxn'], 910)
        # Purchase is 12*(20+2), excluding packaging.
        self.assertEqual(row['planned_purchase_mxn'], 264)
        self.assertEqual(result['cash_days'][0]['planned_closing_mxn'], 736)
        self.assertEqual(result['budget']['status'], 'DENTRO_DEL_ESCENARIO')

    def test_positive_need_obeys_moq_before_pack_rounding(self):
        data = scenario()
        data['stock'][0]['on_hand'] = 15  # available=13, need=1; MOQ=10 => 12
        self.assertEqual(cr.evaluate(data)['skus'][0]['suggested_qty'], 12)
        data['stock'][0]['on_hand'] = 16  # available=14, no purchase despite MOQ
        self.assertEqual(cr.evaluate(data)['skus'][0]['suggested_qty'], 0)

    def test_exact_integer_need_does_not_gain_a_floating_point_unit(self):
        data = scenario()
        data['config'].update(observation_days=25, safety_days=0)
        data['catalog'][0].update(lead_days=100, moq=1, pack_multiple=1)
        data['sales'][0]['delivered_units'] = 7
        data['stock'][0].update(on_hand=0, reserved=0, available_days=25)
        # 7/25 * 100 = exactly 28, despite binary floats representing .28 imprecisely.
        self.assertEqual(cr.evaluate(data)['skus'][0]['suggested_qty'], 28)

    def test_target_price_rounds_up_to_cent_and_infeasible_is_unknown(self):
        data = scenario()
        data['config']['target_margin'] = .3  # 25/.6 = 41.666... => 41.67
        self.assertEqual(cr.evaluate(data)['skus'][0]['target_price_mxn'], 41.67)
        for margin in (.9, .95):
            with self.subTest(margin=margin):
                data['config']['target_margin'] = margin
                row = cr.evaluate(data)['skus'][0]
                self.assertIsNone(row['target_price_mxn'])
                self.assertIsNone(row['suggested_qty'])
                self.assertIn('MARGEN_OBJETIVO_INVIABLE', row['reasons'])

    def test_missing_cost_stays_unknown_while_zero_is_explicit(self):
        data = scenario()
        data['catalog'][0]['other_variable'] = None
        row = cr.evaluate(data)['skus'][0]
        self.assertIsNone(row['unit_variable_cost_mxn'])
        self.assertIsNone(row['target_price_mxn'])
        self.assertIsNone(row['suggested_qty'])
        data['catalog'][0]['other_variable'] = 0
        self.assertEqual(cr.evaluate(data)['skus'][0]['suggested_qty'], 12)

    def test_complete_no_sales_is_zero_incomplete_no_sales_is_unknown(self):
        data = scenario()
        data['sales'] = []
        row = cr.evaluate(data)['skus'][0]
        for field in ('delivered_units', 'net_revenue_mxn', 'variable_cost_mxn', 'contribution_mxn', 'suggested_qty'):
            self.assertEqual(row[field], 0, field)
        data['config']['sales_complete'] = 'NO'
        row = cr.evaluate(data)['skus'][0]
        for field in ('delivered_units', 'net_revenue_mxn', 'variable_cost_mxn', 'contribution_mxn', 'suggested_qty'):
            self.assertIsNone(row[field], field)

    def test_sales_window_is_inclusive_and_excludes_previous_day(self):
        data = scenario()
        row = deepcopy(data['sales'][0])
        row.update(_row=7, date=date(2026, 9, 14), delivered_units=7)
        data['sales'].append(row)
        row = deepcopy(row)
        row.update(_row=8, date=date(2026, 9, 13), delivered_units=999)
        data['sales'].append(row)
        report = cr.evaluate(data)
        self.assertEqual(report['skus'][0]['delivered_units'], 21)
        self.assertEqual(report['skus'][0]['daily_consumption'], 3)

    def test_stale_count_partial_availability_policy_and_price_withhold(self):
        cases = [('stock', 'count_date', date(2026, 9, 19), 'ACTUALIZAR_CONTEO'),
                 ('stock', 'available_days', 6, 'DISPONIBILIDAD_NO_OBSERVADA_COMPLETA'),
                 ('catalog', 'list_price', 30, 'REVISAR_PRECIO_O_COSTO'),
                 ('catalog', 'status', 'PAUSADO', 'PRODUCTO_NO_ACTIVO')]
        for table, field, value, reason in cases:
            with self.subTest(field=field):
                data = scenario()
                data[table][0][field] = value
                row = cr.evaluate(data)['skus'][0]
                self.assertIsNone(row['suggested_qty'])
                self.assertIn(reason, row['reasons'])
        data = scenario()
        data['config']['policy_reviewed'] = 'NO'
        self.assertIsNone(cr.evaluate(data)['skus'][0]['suggested_qty'])

    def test_transit_only_confirmed_within_lead_is_eligible(self):
        data = scenario()
        data['stock'][0].update(in_transit=10, eta=date(2026, 9, 25), transit_confirmed='SI')
        row = cr.evaluate(data)['skus'][0]
        self.assertEqual(row['eligible_transit_units'], 10)
        self.assertEqual(row['suggested_qty'], 0)
        data['stock'][0]['transit_confirmed'] = 'NO'
        row = cr.evaluate(data)['skus'][0]
        self.assertEqual(row['eligible_transit_units'], 0)
        self.assertIsNone(row['suggested_qty'])

    def test_late_or_missing_transit_requires_review(self):
        for eta in (date(2026, 9, 26), date(2026, 9, 20), None):
            with self.subTest(eta=eta):
                data = scenario()
                data['stock'][0].update(in_transit=10, eta=eta, transit_confirmed='SI')
                row = cr.evaluate(data)['skus'][0]
                self.assertEqual(row['eligible_transit_units'], 0)
                self.assertEqual(row['status'], 'REVISAR')
                self.assertIsNone(row['suggested_qty'])

    def test_prior_period_return_restocks_and_cash_are_separate(self):
        data = scenario()
        data['sales'][0].update(delivered_units=1, returned_units=3, restocked_units=2,
                                net_revenue=-100, variable_cost=None)
        result = cr.evaluate(data)
        row = result['skus'][0]
        self.assertFalse(any(i['severity'] == 'ERROR' for i in result['issues']))
        self.assertEqual(row['daily_consumption'], 0)
        self.assertIsNone(row['variable_cost_mxn'])
        self.assertIsNone(row['contribution_mxn'])
        self.assertEqual(result['cash_days'][0]['base_closing_mxn'], 1000)

    def test_partial_historical_cost_is_not_backfilled_from_catalog(self):
        data = scenario()
        second = deepcopy(data['sales'][0])
        second.update(_row=7, date=date(2026, 9, 19), variable_cost=None)
        data['sales'].append(second)
        row = cr.evaluate(data)['skus'][0]
        self.assertEqual(row['net_revenue_mxn'], 2800)
        self.assertIsNone(row['variable_cost_mxn'])
        self.assertIsNone(row['contribution_mxn'])

    def test_monday_shortfall_survives_friday_recovery(self):
        data = scenario()
        data['cash'] = [dict(_row=6, date=date(2026, 9, 21), direction='SALIDA', amount=850),
                        dict(_row=7, date=date(2026, 9, 25), direction='ENTRADA', amount=1000)]
        result = cr.evaluate(data)
        week = result['cash_weeks'][0]
        self.assertEqual(week['base_closing_mxn'], 1150)
        self.assertEqual(week['min_daily_base_mxn'], 150)
        self.assertEqual(week['min_daily_after_mxn'], 150)
        self.assertTrue(week['below_floor'])
        self.assertEqual(result['budget']['available_before_proposals_mxn'], 0)
        self.assertEqual(result['budget']['status'], 'EXCEDE_LIMITE')

    def test_purchase_cash_hits_day_eight_week_two_and_cap(self):
        data = scenario()
        data['stock'][0].update(planned_qty=12, payment_date=date(2026, 9, 28))
        data['config']['purchase_cap'] = 250
        result = cr.evaluate(data)
        self.assertEqual(len(result['cash_days']), 91)
        self.assertEqual(len(result['cash_weeks']), 13)
        self.assertEqual(result['cash_weeks'][0]['planned_purchases_mxn'], 0)
        self.assertEqual(result['cash_weeks'][1]['planned_purchases_mxn'], 264)
        self.assertEqual(result['budget']['status'], 'EXCEDE_LIMITE')

    def test_unknown_plan_does_not_become_zero_and_opening_floor_counts(self):
        data = scenario()
        data['stock'][0]['planned_qty'] = None
        result = cr.evaluate(data)
        self.assertIsNone(result['budget']['planned_purchase_mxn'])
        self.assertIsNone(result['cash_weeks'][0]['planned_closing_mxn'])
        data = scenario()
        data['config']['opening_cash'] = 100
        data['cash'] = [dict(_row=6, date=date(2026, 9, 21), direction='ENTRADA', amount=1000)]
        result = cr.evaluate(data)
        self.assertEqual(result['budget']['min_daily_cash_after_mxn'], 100)
        self.assertEqual(result['budget']['status'], 'EXCEDE_LIMITE')
        data['config']['cash_plan_complete'] = 'NO'
        result = cr.evaluate(data)
        self.assertEqual(result['cash_weeks'], [])
        self.assertIsNone(result['budget']['available_before_proposals_mxn'])

    def test_extreme_values_and_subcent_money_fail_before_arithmetic(self):
        cases = [('catalog', 'purchase_cost', 1e30), ('catalog', 'lead_days', 1e20),
                 ('catalog', 'purchase_cost', .001), ('catalog', 'list_price', 1000000000.01),
                 ('stock', 'planned_qty', 1000001), ('sales', 'delivered_units', 1000001),
                 ('sales', 'net_revenue', -1000000000.01), ('sales', 'variable_cost', .001),
                 ('config', 'safety_days', 3651), ('config', 'as_of', date(9999,12,31)),
                 ('config', 'as_of', date(1899,12,31)), ('config', 'opening_cash', .001)]
        for table, key, value in cases:
            with self.subTest(table=table, key=key, value=value):
                data = scenario()
                (data[table] if table == 'config' else data[table][0])[key] = value
                if key == 'lead_days':
                    data['stock'][0].update(in_transit=1, eta=date(2026,9,21), transit_confirmed='SI')
                report = cr.evaluate(data)
                self.assertEqual(report['status'], 'CORREGIR_CAPTURA')
                self.assertEqual(report['skus'], [])
                self.assertTrue(any(i['severity'] == 'ERROR' for i in report['issues']))

    def test_declared_numeric_boundaries_remain_usable(self):
        data = scenario()
        data['catalog'][0]['lead_days'] = 3650
        data['config']['safety_days'] = 3650
        data['config']['opening_cash'] = 1000000000
        data['stock'][0]['on_hand'] = 1000000
        data['sales'][0]['net_revenue'] = -1000000000
        report = cr.evaluate(data)
        self.assertFalse(any(i['severity'] == 'ERROR' for i in report['issues']))
        self.assertEqual(report['skus'][0]['suggested_qty'], 0)

    def test_negative_floor_cap_and_fractional_windows_rejected_but_negative_opening_allowed(self):
        for key, value in [('cash_floor', -1), ('purchase_cap', -1),
                           ('observation_days', 7.5), ('safety_days', .5)]:
            with self.subTest(key=key):
                data = scenario()
                data['config'][key] = value
                self.assertEqual(cr.evaluate(data)['status'], 'CORREGIR_CAPTURA')
        data = scenario()
        data['config']['opening_cash'] = -100
        report = cr.evaluate(data)
        self.assertFalse(any(i['severity'] == 'ERROR' for i in report['issues']))
        self.assertEqual(report['budget']['min_daily_cash_after_mxn'], -100)
        self.assertEqual(report['budget']['status'], 'EXCEDE_LIMITE')


class ClientWorkbookInputTests(unittest.TestCase):
    def setUp(self):
        root = cr.ROOT / '.local/client-v3-review'
        root.mkdir(parents=True, exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(dir=root)
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.book = self.root / 'synthetic.xlsx'
        write_inputs(self.book, scenario())

    def change(self, sheet, cell, value):
        wb = load_workbook(self.book)
        wb[sheet][cell] = value
        wb.save(self.book)
        wb.close()

    def report(self):
        data, issues = cr.read_inputs(self.book)
        return cr.evaluate(data, issues)

    def test_actual_xlsx_preserves_explicit_zero_and_source_hash(self):
        data, issues = cr.read_inputs(self.book)
        self.assertEqual(issues, [])
        self.assertEqual(data['catalog'][0]['other_variable'], 0)
        self.assertIsNone(data['stock'][0]['eta'])
        self.assertEqual(data['metadata']['source_sha256'], sha256(self.book.read_bytes()).hexdigest())
        self.assertEqual(cr.evaluate(data, issues)['skus'][0]['suggested_qty'], 12)

    def test_formula_or_error_in_input_rejected_without_execution(self):
        for value in ('=1+1', '=HYPERLINK("https://example.invalid", "test")', '#DIV/0!'):
            with self.subTest(value=value):
                self.change('CATALOGO', 'H6', value)
                report = self.report()
                self.assertEqual(report['status'], 'CORREGIR_CAPTURA')
                self.assertIn('INPUT_FORMULA', {i['code'] for i in report['issues']})

    def test_macro_and_external_link_packages_rejected(self):
        for entry in ('xl/vbaProject.bin', 'xl/externalLinks/externalLink1.xml'):
            with self.subTest(entry=entry):
                write_inputs(self.book, scenario())
                with zipfile.ZipFile(self.book, 'a') as archive:
                    archive.writestr(entry, b'synthetic invalid member')
                with self.assertRaisesRegex(ValueError, 'macros|vínculos'):
                    cr.read_inputs(self.book)

    def test_duplicate_catalog_stock_and_day_sku_rejected(self):
        for table, code in [('catalog', 'DUPLICATE_SKU'), ('stock', 'STOCK_SKU'), ('sales', 'DUPLICATE_SALES_AGGREGATE')]:
            with self.subTest(table=table):
                data = scenario()
                row = deepcopy(data[table][0])
                row['_row'] = 7
                data[table].append(row)
                write_inputs(self.book, data)
                report = self.report()
                self.assertEqual(report['status'], 'CORREGIR_CAPTURA')
                self.assertIn(code, {i['code'] for i in report['issues']})

    def test_invalid_pasted_values_and_capacity_rejected(self):
        cases = [('STOCK','C6',-1,'STOCK_RANGE'), ('STOCK','D6',99,'RESERVED_EXCEEDS_STOCK'),
                 ('STOCK','H6',8,'EXPOSURE_RANGE'), ('CATALOGO','O6',0,'CATALOG_RANGE'),
                 ('CATALOGO','K6',1,'CATALOG_RANGE'), ('VENTAS','C6',1.5,'INVALID_NUMBER'),
                 ('VENTAS','E6',1,'RESTOCK_EXCEEDS_RETURN'), ('VENTAS','C6',None,'SALES_VALUE'),
                 ('VENTAS','B6','NO-SUCH-SKU','SALES_REFERENCE'),
                 ('CONFIG','B6',0,'INVALID_PARAMETER'), ('CONFIG','B8',1,'INVALID_PARAMETER'),
                 ('CONFIG','B12','MAYBE','INVALID_CONFIRMATION'),
                 ('STOCK','A106','SYN-OVERFLOW','CAPACITY'),
                 ('VENTAS','H6','synthetic@example.invalid','PRIVATE_REFERENCE'),
                 ('CATALOGO','H6','NaN','INVALID_NUMBER'), ('CATALOGO','H6',True,'INVALID_NUMBER')]
        for sheet, cell, value, code in cases:
            with self.subTest(sheet=sheet, cell=cell, value=value):
                write_inputs(self.book, scenario())
                self.change(sheet, cell, value)
                report = self.report()
                self.assertEqual(report['status'], 'CORREGIR_CAPTURA')
                self.assertIn(code, {i['code'] for i in report['issues']})

    def test_cash_and_sales_date_boundaries_rejected(self):
        data = scenario()
        data['sales'][0]['date'] = date(2026, 9, 21)
        self.assertIn('SALES_REFERENCE', {i['code'] for i in cr.evaluate(data)['issues']})
        for offset in (0, 92):
            data = scenario()
            data['cash'] = [dict(_row=6, date=date(2026,9,20)+timedelta(days=offset), direction='SALIDA', amount=10)]
            self.assertIn('CASH_DATE', {i['code'] for i in cr.evaluate(data)['issues']})

    def test_private_output_boundary_and_no_overwrite(self):
        # Isolate ROOT only for output writes; fixtures remain under the owned review directory.
        with patch.object(cr, 'ROOT', self.root):
            allowed = self.root / '.local/client-runs'
            for path in (self.root/'public-report', allowed, allowed/'../escape'):
                with self.subTest(path=path):
                    with self.assertRaises(ValueError):
                        cr.review_file(self.book, path)
            output = allowed/'synthetic-cut'
            report = cr.review_file(self.book, output)
            self.assertEqual(report['external_execution'], 'PROHIBITED')
            self.assertEqual({p.name for p in output.iterdir()}, {'informe.json','informe.html','correcciones.csv','PARA_EL_ANALISTA.md'})
            with self.assertRaisesRegex(ValueError, 'ya contiene datos'):
                cr.review_file(self.book, output)

    def test_private_html_escapes_customer_text_and_binds_report_hash(self):
        self.change('CATALOGO', 'D6', '<script>alert("synthetic")</script>')
        with patch.object(cr, 'ROOT', self.root):
            output = self.root / '.local/client-runs/escaped'
            cr.review_file(self.book, output)
        html = (output/'informe.html').read_text(encoding='utf-8')
        self.assertNotIn('<script>', html)
        self.assertIn('&lt;script&gt;', html)
        digest = sha256((output/'informe.json').read_bytes()).hexdigest()
        self.assertIn(digest, html)
        self.assertIn(digest, (output/'PARA_EL_ANALISTA.md').read_text(encoding='utf-8'))

    def test_publication_rejects_ungenerated_private_workbook_content(self):
        spec = importlib.util.spec_from_file_location('review_package', cr.ROOT/'scripts/package_client.py')
        package = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(package)
        expected = self.root/'trusted-generated.xlsx'
        expected.write_bytes(self.book.read_bytes())
        package.assert_generated_content(self.book, expected)
        for kind in ('hidden_sheet', 'extra_cell', 'comment', 'hyperlink', 'creator'):
            with self.subTest(kind=kind):
                self.book.write_bytes(expected.read_bytes())
                wb = load_workbook(self.book)
                if kind == 'hidden_sheet':
                    ws = wb.create_sheet('SYNTHETIC_PRIVATE_NOTES')
                    ws.sheet_state = 'hidden'
                    ws['A1'] = 'synthetic@example.invalid'
                elif kind == 'extra_cell':
                    wb['CONFIG']['Z3'] = 'synthetic@example.invalid'
                elif kind == 'comment':
                    wb['CATALOGO']['A6'].comment = Comment('synthetic@example.invalid', 'Synthetic')
                elif kind == 'hyperlink':
                    wb['CATALOGO']['A6'].hyperlink = 'mailto:synthetic@example.invalid'
                else:
                    wb.properties.creator = 'Synthetic Private Customer'
                wb.save(self.book)
                wb.close()
                with self.assertRaises(ValueError):
                    package.assert_generated_content(self.book, expected)

    def test_office_metadata_normalization_preserves_every_other_archive_byte(self):
        spec = importlib.util.spec_from_file_location('review_normalizer', cr.ROOT/'scripts/normalize_client_metadata.py')
        normalizer = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(normalizer)
        with zipfile.ZipFile(self.book) as archive:
            before = {name: archive.read(name) for name in archive.namelist()}
        normalizer.normalize(self.book)
        with zipfile.ZipFile(self.book) as archive:
            after = {name: archive.read(name) for name in archive.namelist()}
        self.assertEqual(set(before), set(after))
        self.assertEqual([name for name in before if before[name] != after[name]], ['docProps/core.xml'])
        wb = load_workbook(self.book)
        self.assertEqual(wb.properties.creator, 'Alma de Lujo')
        self.assertEqual(wb.properties.lastModifiedBy, 'Alma de Lujo')
        wb.close()


if __name__ == '__main__':
    unittest.main()
