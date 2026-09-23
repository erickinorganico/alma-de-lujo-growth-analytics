"""Contract tests for the operating-v1 workbook editing boundary."""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from alma.operating_contracts import CONTRACT_VERSION, SOURCE_NAMES, columns_for
from alma.operating_interchange import parse_pack
from alma.operating_mart_contracts import load_policy
from scripts.operating_pack import synthetic_rows


ROOT = Path(__file__).resolve().parents[1]


class OperatingWorkbookGenerationTests(unittest.TestCase):
    def _build(self, root: Path) -> dict:
        from scripts.build_operating_workbooks import build_operating_workbooks

        return build_operating_workbooks(root)

    def test_blank_and_synthetic_books_follow_the_final_registry(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = self._build(root)
            self.assertEqual(CONTRACT_VERSION, manifest["contract_version"])
            self.assertEqual(22, manifest["relation_count"])
            self.assertEqual(list(SOURCE_NAMES), manifest["relations"])
            self.assertEqual(2, len(manifest["policies"]))
            for kind in ("blank", "synthetic"):
                book_path = root / manifest["workbooks"][kind]["path"]
                pack_path = root / manifest["source_packs"][kind]["path"]
                self.assertEqual(
                    manifest["workbooks"][kind]["sha256"],
                    hashlib.sha256(book_path.read_bytes()).hexdigest(),
                )
                workbook = load_workbook(book_path, data_only=False, keep_links=False)
                self.assertEqual(
                    ["INICIO", "DICCIONARIO", "COMPLETITUD", *SOURCE_NAMES],
                    workbook.sheetnames,
                )
                for relation in SOURCE_NAMES:
                    with self.subTest(kind=kind, relation=relation):
                        sheet = workbook[relation]
                        self.assertEqual(
                            list(columns_for(relation)),
                            [sheet.cell(6, column).value for column in range(1, len(columns_for(relation)) + 1)],
                        )
                        self.assertEqual("A7", sheet.freeze_panes)
                        self.assertEqual("landscape", sheet.page_setup.orientation)
                        self.assertEqual("$1:$6", sheet.print_title_rows)
                        self.assertTrue(sheet.sheet_properties.pageSetUpPr.fitToPage)
                        self.assertIn(sheet.page_setup.fitToWidth, {None, 1})
                        self.assertTrue(sheet.auto_filter.ref.startswith("A6:"))
                        self.assertTrue(sheet.data_validations.count >= 1)
                        rows = [
                            [sheet.cell(row, column).value for column in range(1, len(columns_for(relation)) + 1)]
                            for row in range(7, sheet.max_row + 1)
                        ]
                        rows = [row for row in rows if any(value is not None for value in row)]
                        if kind == "blank":
                            self.assertEqual([], rows)
                        else:
                            self.assertEqual(len(synthetic_rows()[relation]), len(rows))
                            self.assertIn("SYNTHETIC", str(sheet["A2"].value).upper())
                workbook.close()
                if kind == "synthetic":
                    parsed = parse_pack(pack_path)
                    self.assertEqual(list(SOURCE_NAMES), list(parsed["tables"]))
                else:
                    metadata = json.loads((pack_path / "metadata.json").read_text(encoding="utf-8"))
                    self.assertEqual("BLANK", metadata["input_class"])

    def test_completeness_matrix_covers_each_relation_and_interactions(self) -> None:
        from scripts.build_operating_workbooks import COMPLETENESS_MATRIX, omission_impact

        self.assertEqual(set(SOURCE_NAMES), set(COMPLETENESS_MATRIX))
        for relation in SOURCE_NAMES:
            with self.subTest(relation=relation):
                row = COMPLETENESS_MATRIX[relation]
                self.assertIn(row["classification"], {"REQUIRED", "OPTIONAL"})
                self.assertIn(row["omission_state"], {"UNKNOWN", "PARTIAL", "REVIEW", "BLOCKED"})
                self.assertTrue(row["metric_families"])
                self.assertEqual(row["omission_state"], omission_impact((relation,))["state"])
        self.assertEqual("BLOCKED", omission_impact(("sku_catalog", "sales_aggregates"))["state"])
        self.assertEqual("UNKNOWN", omission_impact(("sales_aggregates", "quality_events"))["state"])
        self.assertEqual(
            sorted(set(COMPLETENESS_MATRIX["sales_aggregates"]["metric_families"])
                       | set(COMPLETENESS_MATRIX["quality_events"]["metric_families"])),
            omission_impact(("sales_aggregates", "quality_events"))["metric_families"],
        )

    def test_policy_materials_are_separate_hash_bound_and_unapproved(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = self._build(root)
            self.assertNotIn("policies", manifest["relations"])
            for policy in manifest["policies"]:
                path = root / policy["path"]
                content = json.loads(path.read_text(encoding="utf-8"))
                self.assertEqual(policy["sha256"], hashlib.sha256(path.read_bytes()).hexdigest())
                loaded = load_policy(path, as_of="2026-09-21", real_cut=False)
                self.assertEqual(content["status"], loaded.status)
                self.assertIsNone(content["owner_approval_ref"])
            self.assertEqual(
                {"REVIEW", "SYNTHETIC_EXAMPLE"},
                {json.loads((root / policy["path"]).read_text(encoding="utf-8"))["status"]
                 for policy in manifest["policies"]},
            )

    def test_generated_manifest_rehashes_every_file(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            manifest = self._build(root)
            persisted = json.loads((root / "workbook-manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(manifest, persisted)
            for relative, digest in manifest["files"].items():
                self.assertEqual(digest, hashlib.sha256((root / relative).read_bytes()).hexdigest())


class WorkbookImportTests(unittest.TestCase):
    def _generated(self, root: Path) -> tuple[Path, Path]:
        from scripts.build_operating_workbooks import build_operating_workbooks

        build_operating_workbooks(root / "generated")
        return (
            root / "generated" / "operating-v1-synthetic.xlsx",
            root / "generated" / "source-packs" / "synthetic",
        )

    def _export(self, workbook: Path, root: Path) -> dict:
        from alma.operating_workbook import export_workbook_to_pack

        private_root = root / "private"
        private_root.mkdir(exist_ok=True)
        return export_workbook_to_pack(workbook, private_root / "export", private_root=private_root)

    @staticmethod
    def _raw_workbook_rows(path: Path, relation: str) -> list[list[object]]:
        workbook = load_workbook(path, data_only=False, keep_links=False)
        sheet = workbook[relation]
        width = len(columns_for(relation))
        rows = [
            [sheet.cell(row, column).value for column in range(1, width + 1)]
            for row in range(7, sheet.max_row + 1)
        ]
        workbook.close()
        return [row for row in rows if any(value is not None for value in row)]

    def test_workbook_pack_manifest_oracle_covers_all_relations(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            workbook_path, _ = self._generated(root)
            original = workbook_path.read_bytes()
            receipt = self._export(workbook_path, root)
            self.assertEqual("PASS", receipt["status"])
            self.assertEqual(list(SOURCE_NAMES), receipt["relations"])
            self.assertEqual(hashlib.sha256(original).hexdigest(), receipt["workbook_sha256"])
            self.assertEqual(original, workbook_path.read_bytes())
            pack = Path(receipt["pack_path"])
            manifest = json.loads(Path(receipt["manifest_path"]).read_text(encoding="utf-8"))
            parsed = parse_pack(pack)
            self.assertEqual("SYNTHETIC_EXAMPLE", parsed["metadata"]["input_class"])
            self.assertEqual(22, len(manifest["relations"]))
            for relation in SOURCE_NAMES:
                with self.subTest(relation=relation):
                    fields = list(columns_for(relation))
                    workbook_rows = self._raw_workbook_rows(workbook_path, relation)
                    with (pack / f"{relation}.csv").open(encoding="utf-8", newline="") as stream:
                        csv_rows = list(csv.DictReader(stream))
                    self.assertEqual(len(workbook_rows), len(csv_rows))
                    self.assertEqual(fields, list(csv_rows[0]) if csv_rows else fields)
                    self.assertEqual(
                        hashlib.sha256((pack / f"{relation}.csv").read_bytes()).hexdigest(),
                        manifest["relations"][relation]["sha256"],
                    )
                    for raw, exported in zip(workbook_rows, csv_rows, strict=True):
                        expected = ["" if value is None else value.isoformat() if isinstance(value, (date, datetime)) else str(value)
                                    for value in raw]
                        self.assertEqual(expected, [exported[field] for field in fields])
                    cents = [field for field in fields if field.endswith("_cents")]
                    for field in cents:
                        self.assertEqual("MXN cents", manifest["relations"][relation]["units"][field])
                        for row in csv_rows:
                            if row[field] != "":
                                self.assertEqual(Decimal(row[field]), Decimal(row[field]).to_integral_value())
            self.assertIsNone(parsed["tables"]["loans"][0]["returned_date"])
            self.assertEqual(0, parsed["tables"]["purchase_receipts"][0]["rejected_units"])

    def test_structural_and_hidden_content_fail_before_publication(self) -> None:
        cases = (
            ("extra_sheet", lambda wb: wb.create_sheet("PRIVATE_NOTES"), "workbook.sheet_set"),
            ("missing_sheet", lambda wb: wb.remove(wb["loans"]), "workbook.sheet_set"),
            ("header", lambda wb: setattr(wb["sales_aggregates"]["A6"], "value", "wrong"), "workbook.header"),
            ("formula", lambda wb: setattr(wb["sales_aggregates"]["H7"], "value", "=1+1"), "workbook.formula"),
            ("hidden_sheet", lambda wb: setattr(wb["loans"], "sheet_state", "hidden"), "workbook.hidden_sheet"),
            ("hidden_row", lambda wb: setattr(wb["sales_aggregates"].row_dimensions[7], "hidden", True), "workbook.hidden_row"),
            ("hidden_column", lambda wb: setattr(wb["sales_aggregates"].column_dimensions["A"], "hidden", True), "workbook.hidden_column"),
            ("comment", lambda wb: setattr(wb["sales_aggregates"]["A7"], "comment", __import__("openpyxl").comments.Comment("private", "x")), "workbook.comment"),
            ("external_link", lambda wb: setattr(wb["sales_aggregates"]["A7"], "hyperlink", "https://example.com"), "workbook.hyperlink"),
        )
        for name, mutate, code in cases:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source, _ = self._generated(root)
                candidate = root / f"{name}.xlsx"
                shutil.copyfile(source, candidate)
                workbook = load_workbook(candidate, data_only=False, keep_links=False)
                mutate(workbook)
                workbook.save(candidate)
                workbook.close()
                from alma.operating_workbook import WorkbookContractError, export_workbook_to_pack
                private = root / "private"
                private.mkdir()
                with self.assertRaises(WorkbookContractError) as raised:
                    export_workbook_to_pack(candidate, private / "export", private_root=private)
                self.assertEqual(code, raised.exception.code)
                self.assertFalse((private / "export").exists())

    def test_value_relationship_pii_and_overwrite_controls(self) -> None:
        cases = (
            ("fractional_cents", "sales_aggregates", "H7", 10.5, "value.integer"),
            ("bad_date", "sales_aggregates", "A7", "2026-99-99", "value.date"),
            ("duplicate_key", "sku_catalog", "A8", "synthetic:sku-001", "value.required"),
            ("broken_relation", "sales_aggregates", "B7", "synthetic:missing", "relation.sku"),
            ("pii", "sales_aggregates", "K7", "5551234567", "value.pii"),
        )
        for name, relation, cell, value, expected in cases:
            with self.subTest(name=name), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                source, _ = self._generated(root)
                candidate = root / f"{name}.xlsx"
                shutil.copyfile(source, candidate)
                workbook = load_workbook(candidate, data_only=False, keep_links=False)
                if name == "duplicate_key":
                    for column in range(1, len(columns_for(relation)) + 1):
                        workbook[relation].cell(8, column).value = workbook[relation].cell(7, column).value
                else:
                    workbook[relation][cell] = value
                workbook.save(candidate)
                workbook.close()
                from alma.operating_workbook import WorkbookContractError, export_workbook_to_pack
                private = root / "private"
                private.mkdir()
                with self.assertRaises(WorkbookContractError) as raised:
                    export_workbook_to_pack(candidate, private / "export", private_root=private)
                self.assertEqual(expected, raised.exception.upstream_code or raised.exception.code)
                self.assertTrue(raised.exception.sheet)
                self.assertTrue(raised.exception.correction)
                self.assertFalse((private / "export").exists())

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, _ = self._generated(root)
            first = self._export(source, root)
            from alma.operating_workbook import WorkbookContractError, export_workbook_to_pack
            with self.assertRaises(WorkbookContractError) as raised:
                export_workbook_to_pack(source, first["bundle_path"], private_root=root / "private")
            self.assertEqual("output.exists", raised.exception.code)

    def test_archive_member_injection_fails_closed(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source, _ = self._generated(root)
            candidate = root / "injected.xlsx"
            shutil.copyfile(source, candidate)
            with zipfile.ZipFile(candidate, "a") as archive:
                archive.writestr("xl/externalLinks/externalLink1.xml", "private")
            from alma.operating_workbook import WorkbookContractError, export_workbook_to_pack
            private = root / "private"
            private.mkdir()
            with self.assertRaises(WorkbookContractError) as raised:
                export_workbook_to_pack(candidate, private / "export", private_root=private)
            self.assertEqual("workbook.external_content", raised.exception.code)
            self.assertFalse((private / "export").exists())

    pass


if __name__ == "__main__":
    unittest.main()
