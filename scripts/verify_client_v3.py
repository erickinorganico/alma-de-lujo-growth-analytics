"""Prepare synthetic Excel probes; compare real-engine cached cells with Decimal oracle.

Preparation deliberately does not supply formula caches. Run the separate Excel
recalculation script before check. CI rechecks the exact delivered workbooks and
their recorded Excel-engine hashes, without pretending to run Excel on Linux.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from datetime import date, timedelta
from hashlib import sha256
import importlib.util
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from alma.client_review import CONTRACT, read_inputs, evaluate
from openpyxl import load_workbook


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def prepare(directory):
    builder = load_module('client_builder', ROOT/'scripts/build_client_workbook.py')
    test = load_module('independent_probe', ROOT/'tests/test_client_review.py')
    directory.mkdir(parents=True, exist_ok=True)
    baseline = test.scenario()
    cases = {'hand_calculated': deepcopy(baseline)}
    cases['hand_calculated']['stock'][0].update(planned_qty=12, payment_date=date(2026,9,21))
    for name in ['daily_shortfall','missing_cost','incomplete_sales','blank_purchase','late_transit','overdue_transit','invalid_pack','price_below_target','infeasible_margin','unknown_bank','partial_exposure','prior_period_return','zero_sales','invalid_paste','duplicate_sales','missing_stock']:
        cases[name] = deepcopy(baseline)
    cases['daily_shortfall']['cash'] = [dict(_row=6,date=date(2026,9,21),direction='SALIDA',category='Ejemplo',amount=850,source_ref='SYN-OUT'),dict(_row=7,date=date(2026,9,25),direction='ENTRADA',category='Ejemplo',amount=1000,source_ref='SYN-IN')]
    cases['missing_cost']['catalog'][0]['packaging']=None
    cases['incomplete_sales']['config']['sales_complete']='NO'
    cases['blank_purchase']['stock'][0]['planned_qty']=None
    cases['late_transit']['stock'][0].update(in_transit=12,eta=date(2026,10,10),transit_confirmed='SI')
    cases['overdue_transit']['stock'][0].update(in_transit=12,eta=date(2026,9,19),transit_confirmed='SI')
    cases['invalid_pack']['stock'][0].update(planned_qty=11,payment_date=date(2026,9,21))
    cases['price_below_target']['catalog'][0]['list_price']=30
    cases['infeasible_margin']['config']['target_margin']=.95
    cases['unknown_bank']['config']['opening_cash']=None
    cases['partial_exposure']['stock'][0]['available_days']=4
    cases['prior_period_return']['sales'][0].update(delivered_units=0,returned_units=2,restocked_units=1,net_revenue=-200,variable_cost=-20)
    cases['zero_sales']['sales']=[]
    cases['invalid_paste']['stock'][0]['reserved']=100
    cases['duplicate_sales']['sales'].append(deepcopy(cases['duplicate_sales']['sales'][0]))
    cases['duplicate_sales']['sales'][1]['_row']=7
    cases['missing_stock']['stock']=[]
    cases['quantity_binary_boundary']=deepcopy(baseline)
    cases['quantity_binary_boundary']['config'].update(observation_days=25,safety_days=0)
    cases['quantity_binary_boundary']['catalog'][0].update(lead_days=100,moq=1,pack_multiple=1)
    cases['quantity_binary_boundary']['stock'][0].update(on_hand=0,reserved=0,available_days=25)
    cases['quantity_binary_boundary']['sales'][0].update(delivered_units=7,net_revenue=700,variable_cost=245)
    for name in ['blank_safety','blank_target','blank_commission','opening_breach','duplicate_stock_missing_sku','invalid_text_cost','over_purchase_cap']:
        cases[name]=deepcopy(baseline)
    cases['blank_safety']['config']['safety_days']=None
    cases['blank_target']['config']['target_margin']=None
    cases['blank_commission']['catalog'][0]['commission_rate']=None
    cases['opening_breach']['config']['opening_cash']=100
    cases['opening_breach']['cash']=[dict(_row=6,date=date(2026,9,21),direction='ENTRADA',category='Ejemplo',amount=1000,source_ref='SYN-IN')]
    duplicate=cases['duplicate_stock_missing_sku']
    duplicate['catalog'].append(deepcopy(duplicate['catalog'][0]));duplicate['catalog'][1].update(_row=7,sku='SYN-SECOND')
    duplicate['stock'].append(deepcopy(duplicate['stock'][0]));duplicate['stock'][1]['_row']=7
    cases['invalid_text_cost']['catalog'][0]['purchase_cost']='desconocido'
    cases['over_purchase_cap']['config']['purchase_cap']=100
    cases['over_purchase_cap']['stock'][0].update(planned_qty=12,payment_date=date(2026,9,21))
    for name in ['negative_cash_floor','text_observation_window','fractional_safety']:
        cases[name]=deepcopy(baseline)
    cases['negative_cash_floor']['config']['cash_floor']=-1
    cases['text_observation_window']['config']['observation_days']='siete'
    cases['fractional_safety']['config']['safety_days']=.5
    for name, data in cases.items():
        payload={'contract_version':'3.0','workbook_type':'EJEMPLO_SINTETICO','parameters':data['config']}
        for sheet,key in [('CATALOGO','catalog'),('VENTAS','sales'),('STOCK','stock'),('CAJA','cash')]:payload[sheet]=data[key]
        builder.build_workbook(directory/(name+'.xlsx'),payload,CONTRACT)
    print(json.dumps({'prepared':len(cases),'directory':str(directory)}))


def workbook_check(path, receipt):
    actual_hash=sha256(path.read_bytes()).hexdigest()
    checks=[]
    def check(name, actual, expected):
        a=None if actual=='' else actual;e=None if expected=='' else expected
        equal=abs(a-e)<=0.000001 if isinstance(a,(int,float)) and not isinstance(a,bool) and isinstance(e,(int,float)) and not isinstance(e,bool) else a==e
        checks.append(dict(check=name,actual=a,expected=e,passed=equal))
    check('exact_excel_recalculated_bytes',actual_hash,receipt.get('sha256'))
    check('actual_engine',receipt.get('engine'),'Microsoft Excel')
    check('recalculated',receipt.get('recalculated'),True)
    formulas=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True)
    errors=[];formula_count=0
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.data_type=='f':formula_count+=1
                if cached[ws.title][cell.coordinate].data_type=='e':errors.append(ws.title+'!'+cell.coordinate+':'+str(cached[ws.title][cell.coordinate].value))
    check('formula_errors',errors,[])
    data,issues=read_inputs(path);report=evaluate(data,issues)
    def at(sheet,cell):return cached[sheet][cell].value
    if report['status']=='CORREGIR_CAPTURA':
        inventory_error=any(x.get('severity')=='ERROR' and (str(x.get('cell','')).startswith(('CATALOGO!','VENTAS!','STOCK!')) or str(x.get('cell','')) in {'CONFIG!B5','CONFIG!B6','CONFIG!B7','CONFIG!B8'}) for x in report['issues'])
        if inventory_error:
            for record in data['catalog']:
                check('invalid_capture_withholds_quantity_'+str(record['_row']),at('DECISIONES','M'+str(record['_row'])),None)
        if inventory_error:
            check('invalid_capture_withholds_purchase_total_N6',at('FLUJO_13_SEMANAS','N6'),None)
        check('invalid_capture_withholds_cash_N8',at('FLUJO_13_SEMANAS','N8'),None)
        check('invalid_capture_no_confident_cash_headline',at('FLUJO_13_SEMANAS','N3') in ('PLAN SOBRE PISO','DENTRO DEL ESCENARIO'),False)
    else:
        by_sku={r['sku']:r for r in report['skus']}
        for source in data['catalog']:
            row=source['_row'];r=by_sku[source['sku']]
            for col,key in [('P','unit_variable_cost_mxn'),('Q','unit_contribution_mxn'),('R','unit_margin'),('S','target_price_mxn')]:
                check(f'CATALOGO!{col}{row}',at('CATALOGO',f'{col}{row}'),r[key])
            check(f'DECISIONES!M{row}',at('DECISIONES',f'M{row}'),r['suggested_qty'])
            check(f'DECISIONES!N{row}',at('DECISIONES',f'N{row}'),r['planned_qty'])
        for source in data['stock']:
            row=source['_row'];r=by_sku[source['sku']]
            for col,key in [('L','available_units'),('M','daily_consumption'),('R','planned_purchase_mxn')]:
                check(f'STOCK!{col}{row}',at('STOCK',f'{col}{row}'),r[key])
        b=report['budget']
        headline=str(at('FLUJO_13_SEMANAS','N3') or '')
        if b['status']=='EXCEDE_LIMITE':
            check('cash_headline_breach',any(t in headline for t in ['BRECHA','EXCEDE','REBASA']),True)
        elif b['status']=='DENTRO_DEL_ESCENARIO':
            check('cash_headline_within_scenario',headline,'PLAN SOBRE PISO')
        else:
            check('cash_headline_unknown',any(t in headline for t in ['INCOMPLETO','PENDIENTE','INVALID','REVISAR','SIN BASE']),True)
        for cell,key in [('N5','available_before_proposals_mxn'),('N6','planned_purchase_mxn'),('N8','min_daily_cash_after_mxn')]:
            check('FLUJO_13_SEMANAS!'+cell,at('FLUJO_13_SEMANAS',cell),b[key])
        for week in report['cash_weeks']:
            row=week['week']+5
            for col,key in [('F','base_closing_mxn'),('G','planned_purchases_mxn'),('H','planned_closing_mxn'),('I','min_daily_base_mxn'),('J','min_daily_after_mxn')]:
                check(f'FLUJO_13_SEMANAS!{col}{row}',at('FLUJO_13_SEMANAS',f'{col}{row}'),week[key])
        for day in report['cash_days']:
            row=(date.fromisoformat(day['date'])-date.fromisoformat(report['as_of'])).days+5
            for col,key in [('D','base_closing_mxn'),('F','planned_closing_mxn')]:
                check(f'AUX_CALC_91D!{col}{row}',at('AUX_CALC_91D',f'{col}{row}'),day[key])
    formulas.close();cached.close()
    return dict(file=path.name,sha256=actual_hash,synthetic_or_blank=True,formula_count=formula_count,checks_count=len(checks),passed=all(c['passed'] for c in checks),failures=[c for c in checks if not c['passed']],oracle_status=report['status'])


def check_directory(directory, engine_receipt, output):
    receipts=json.loads(engine_receipt.read_text(encoding='utf-8-sig'))
    if isinstance(receipts,dict):receipts=[receipts]
    by_name={r['file']:r for r in receipts}
    results=[workbook_check(p,by_name.get(p.name,{})) for p in sorted(directory.glob('*.xlsx'))]
    report=dict(version='3.0',passed=bool(results) and all(r['passed'] for r in results),engine_receipt=str(engine_receipt),workbooks=results,limitations='Recorded Excel 16 recalculation, exact-byte hashes and independent Decimal checks. CI verifies recorded results; it does not run Excel. No real client records or adoption measured.')
    output.parent.mkdir(parents=True,exist_ok=True);output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return 0 if report['passed'] else 1


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('prepare');p.add_argument('--directory',type=Path,required=True)
    p=sub.add_parser('check');p.add_argument('--directory',type=Path,required=True);p.add_argument('--engine-receipt',type=Path,required=True);p.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    if args.command=='prepare':prepare(args.directory)
    else:raise SystemExit(check_directory(args.directory,args.engine_receipt,args.output))
