#!/usr/bin/env python3
"""Build contract-driven operating-v1 XLSX books and canonical CSV packs."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import sys
import tempfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Sequence

import xlsxwriter


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from alma.operating_contracts import CONTRACT_VERSION, SOURCE_NAMES, SOURCES, canonical_json  # noqa: E402
from alma.operating_interchange import parse_pack  # noqa: E402
from alma.operating_mart_contracts import load_policy  # noqa: E402
from alma.operating_marts import _definitions  # noqa: E402
from scripts.operating_pack import (  # noqa: E402
    blank_metadata, initialize_pack, synthetic_metadata, synthetic_rows,
)


MAX_INPUT_ROWS = 1000
SUPPORT_SHEETS = ("INICIO", "DICCIONARIO", "COMPLETITUD")
POLICY_FILES = (
    "operating-metrics-review-template-v1.json",
    "operating-metrics-synthetic-v1.json",
)
DOMAINS = {
    "sku_catalog": "CATALOGO",
    "sales_aggregates": "COMERCIO",
    "availability_daily": "COMERCIO",
    "unmet_demand": "COMERCIO",
    "inventory_counts": "INVENTARIO",
    "inventory_movements": "INVENTARIO",
    "inventory_reservations": "INVENTARIO",
    "cost_versions": "COSTOS",
    "cost_components": "COSTOS",
    "cost_allocations": "COSTOS",
    "purchase_orders": "COMPRAS",
    "purchase_receipts": "COMPRAS",
    "obligations": "FINANZAS",
    "obligation_payments": "FINANZAS",
    "cash_events": "CAJA",
    "cash_balance_evidence": "CAJA",
    "budgets": "PRESUPUESTO",
    "budget_allocations": "PRESUPUESTO",
    "expenses": "PRESUPUESTO",
    "quality_events": "CALIDAD",
    "sales_readiness": "PREPARACION",
    "loans": "CUSTODIA",
}


def _metric_family(metric_id: str) -> str:
    for token, family in (
        ("cash", "cash"), ("budget", "budget"), ("obligation", "obligations"),
        ("purchase", "purchases"), ("cost", "costs"), ("inventory", "inventory"),
        ("availability", "availability"), ("stockout", "availability"),
        ("return", "returns_quality"), ("quality", "returns_quality"),
        ("sell", "commerce"), ("variant", "commerce"), ("unmet", "demand"),
    ):
        if token in metric_id:
            return family
    return "operating_metrics"


def _metric_families() -> dict[str, list[str]]:
    by_source = {name: set() for name in SOURCE_NAMES}
    for definition in _definitions().values():
        for source in definition.sources:
            by_source[source].add(_metric_family(definition.id))
    by_source["sku_catalog"].add("identity_and_relationships")
    by_source["sales_readiness"].add("commercial_readiness")
    return {name: sorted(values) for name, values in by_source.items()}


_UNKNOWN = {
    "sales_aggregates", "inventory_counts", "cost_versions",
    "cash_events", "cash_balance_evidence",
}
_REVIEW = {"sales_readiness", "budgets", "budget_allocations"}
COMPLETENESS_MATRIX = {
    name: {
        "relation": name,
        "domain": DOMAINS[name],
        "classification": "REQUIRED" if name == "sku_catalog" else "OPTIONAL",
        "metric_families": _metric_families()[name],
        "omission_state": (
            "BLOCKED" if name == "sku_catalog" else
            "UNKNOWN" if name in _UNKNOWN else
            "REVIEW" if name in _REVIEW else
            "PARTIAL"
        ),
    }
    for name in SOURCE_NAMES
}


def omission_impact(omitted: Sequence[str]) -> dict[str, Any]:
    """Return the declared decision impact for one or more omitted sources."""

    unknown = set(omitted) - set(SOURCE_NAMES)
    if unknown:
        raise ValueError("unknown relation: " + ", ".join(sorted(unknown)))
    priorities = {"PARTIAL": 1, "REVIEW": 2, "UNKNOWN": 3, "BLOCKED": 4}
    rows = [COMPLETENESS_MATRIX[name] for name in omitted]
    state = max((row["omission_state"] for row in rows), key=priorities.get, default="PARTIAL")
    return {
        "omitted": list(omitted),
        "state": state,
        "metric_families": sorted({family for row in rows for family in row["metric_families"]}),
    }


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _formats(workbook: xlsxwriter.Workbook) -> dict[str, Any]:
    return {
        "title": workbook.add_format({"bold": True, "font_size": 18, "font_color": "#FFFFFF", "bg_color": "#183C32", "align": "left", "valign": "vcenter"}),
        "subtitle": workbook.add_format({"font_size": 10, "font_color": "#183C32", "bg_color": "#EDE5D5", "text_wrap": True, "valign": "vcenter"}),
        "synthetic": workbook.add_format({"bold": True, "font_color": "#745323", "bg_color": "#F4E4C5", "text_wrap": True}),
        "header": workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#246B61", "border": 1, "text_wrap": True, "valign": "vcenter"}),
        "input": workbook.add_format({"bg_color": "#DCEEFF", "font_color": "#17231F", "border": 1, "locked": False}),
        "body": workbook.add_format({"font_color": "#17231F", "border": 1, "valign": "top"}),
        "wrap": workbook.add_format({"font_color": "#17231F", "border": 1, "text_wrap": True, "valign": "top"}),
        "link": workbook.add_format({"font_color": "#176E66", "underline": True}),
        "note": workbook.add_format({"font_color": "#394A43", "bg_color": "#F7F3EA", "text_wrap": True, "valign": "top"}),
    }


def _setup_page(sheet: Any, label: str) -> None:
    sheet.set_landscape()
    sheet.fit_to_pages(1, 0)
    sheet.set_paper(9)
    sheet.set_margins(0.3, 0.3, 0.45, 0.45)
    sheet.repeat_rows(0, 5)
    sheet.set_header(f"&LAlma de Lujo | {label}&R{CONTRACT_VERSION}")
    sheet.set_footer("&LArchivo local · datos agregados&C&P / &N&R&F")


def _write_support_sheets(workbook: xlsxwriter.Workbook, kind: str, formats: dict[str, Any]) -> None:
    synthetic = kind == "synthetic"
    start = workbook.add_worksheet("INICIO")
    start.set_tab_color("#B98236" if synthetic else "#246B61")
    start.merge_range("A1:F1", "ALMA OS · CAPTURA OPERATIVA v1", formats["title"])
    marker = ("SYNTHETIC_EXAMPLE · EJEMPLO SINTETICO · TODOS LOS DATOS SON INVENTADOS" if synthetic else
              "PLANTILLA EN BLANCO · VACIO NO SIGNIFICA CERO")
    start.merge_range("A2:F2", marker, formats["synthetic"] if synthetic else formats["subtitle"])
    metadata = [
        ("Contrato", CONTRACT_VERSION),
        ("Clase", "SYNTHETIC_EXAMPLE" if synthetic else "BLANK"),
        ("Corte", "2026-09-21T23:59:59-07:00" if synthetic else "POR_COMPLETAR"),
        ("Zona horaria", "America/Tijuana" if synthetic else "POR_COMPLETAR"),
        ("Moneda", "MXN"),
    ]
    for row, (label, value) in enumerate(metadata, start=3):
        start.write(row - 1, 0, label, formats["header"])
        start.write(row - 1, 1, value, formats["body"])
    start.write(9, 0, "Orden", formats["header"])
    start.write(9, 1, "Dominio", formats["header"])
    start.write(9, 2, "Relación", formats["header"])
    start.write(9, 3, "Clasificación", formats["header"])
    start.write(9, 4, "Si falta", formats["header"])
    start.write(9, 5, "Uso", formats["header"])
    for index, relation in enumerate(SOURCE_NAMES, start=1):
        row = 9 + index
        matrix = COMPLETENESS_MATRIX[relation]
        start.write(row, 0, index, formats["body"])
        start.write(row, 1, matrix["domain"], formats["body"])
        start.write_url(row, 2, f"internal:'{relation}'!A1", formats["link"], relation)
        start.write(row, 3, matrix["classification"], formats["body"])
        start.write(row, 4, matrix["omission_state"], formats["body"])
        start.write(row, 5, ", ".join(matrix["metric_families"]), formats["wrap"])
    start.set_column("A:A", 12)
    start.set_column("B:B", 17)
    start.set_column("C:C", 28)
    start.set_column("D:E", 17)
    start.set_column("F:F", 45)
    start.freeze_panes(10, 0)
    _setup_page(start, marker)

    dictionary = workbook.add_worksheet("DICCIONARIO")
    headers = ("dominio", "relacion", "grano", "campo", "tipo", "unidad", "nullable", "enum", "proveniencia")
    for column, header in enumerate(headers):
        dictionary.write(0, column, header, formats["header"])
    row = 1
    for relation in SOURCE_NAMES:
        contract = SOURCES[relation]
        for field in contract["fields"]:
            values = (
                DOMAINS[relation], relation, contract["grain"], field["name"], field["type"],
                field["unit"], "SI" if field["nullable"] else "NO",
                " | ".join(str(value) for value in (field["enum"] or ())), field["provenance"],
            )
            for column, value in enumerate(values):
                dictionary.write(row, column, value, formats["wrap"])
            row += 1
    dictionary.autofilter(0, 0, row - 1, len(headers) - 1)
    dictionary.freeze_panes(1, 0)
    dictionary.set_column(0, 2, 24)
    dictionary.set_column(3, 8, 22)
    _setup_page(dictionary, marker)

    completeness = workbook.add_worksheet("COMPLETITUD")
    headers = ("relacion", "dominio", "clasificacion", "estado_si_falta", "familias_metricas",
               "coverage_status", "window_start", "window_end", "regla")
    for column, header in enumerate(headers):
        completeness.write(0, column, header, formats["header"])
    metadata = synthetic_metadata() if synthetic else blank_metadata()
    for row, relation in enumerate(SOURCE_NAMES, start=1):
        matrix = COMPLETENESS_MATRIX[relation]
        coverage = metadata["coverage"][relation]
        values = (
            relation, matrix["domain"], matrix["classification"], matrix["omission_state"],
            ", ".join(matrix["metric_families"]),
            coverage["status"], coverage["window_start"], coverage["window_end"],
            "Vacío = desconocido; registra ZERO sólo cuando observaste cero en una ventana declarada.",
        )
        for column, value in enumerate(values):
            completeness.write(row, column, value, formats["wrap"])
    completeness.autofilter(0, 0, len(SOURCE_NAMES), len(headers) - 1)
    completeness.freeze_panes(1, 0)
    completeness.set_column(0, 3, 23)
    completeness.set_column(4, 4, 52)
    completeness.set_column(5, 7, 18)
    completeness.set_column(8, 8, 52)
    _setup_page(completeness, marker)


def _validation(sheet: Any, column: int, field: dict[str, Any]) -> None:
    first, last = 6, 5 + MAX_INPUT_ROWS
    options: dict[str, Any] = {
        "ignore_blank": bool(field["nullable"]),
        "error_title": "Valor fuera del contrato",
        "error_message": f"Corrige {field['name']} según el diccionario {CONTRACT_VERSION}.",
        "input_title": field["name"][:32],
        "input_message": f"{field['type']} · {field['unit']}"[:255],
    }
    if field["enum"]:
        options.update(validate="list", source=list(field["enum"]))
    elif field["type"] == "date":
        options.update(validate="date", criteria="between", minimum=datetime(1900, 1, 1), maximum=datetime(2100, 12, 31))
    elif field["type"] in {"integer", "nonnegative_integer", "signed_integer"}:
        minimum = 0 if field["type"] == "nonnegative_integer" else -2_147_483_648
        options.update(validate="integer", criteria="between", minimum=minimum, maximum=2_147_483_647)
    else:
        options.update(validate="length", criteria="between", minimum=1, maximum=64)
    sheet.data_validation(first, column, last, column, options)


def _write_relation_sheet(workbook: xlsxwriter.Workbook, relation: str, rows: list[dict[str, Any]],
                          kind: str, formats: dict[str, Any]) -> None:
    contract = SOURCES[relation]
    sheet = workbook.add_worksheet(relation)
    sheet.merge_range(0, 0, 0, max(0, len(contract["fields"]) - 1),
                      f"{DOMAINS[relation]} · {relation}", formats["title"])
    marker = "SYNTHETIC_EXAMPLE · EJEMPLO SINTETICO · datos inventados" if kind == "synthetic" else "PLANTILLA EN BLANCO · no completes faltantes con cero"
    sheet.merge_range(1, 0, 1, max(0, len(contract["fields"]) - 1), marker,
                      formats["synthetic"] if kind == "synthetic" else formats["subtitle"])
    note = (
        f"Grano: {contract['grain']} · Clave: {', '.join(contract['primary_key'])} · "
        f"Fuente/proveniencia: {contract['provenance_field']} · Estado al omitir: "
        f"{COMPLETENESS_MATRIX[relation]['omission_state']}."
    )
    sheet.merge_range(2, 0, 3, max(0, len(contract["fields"]) - 1), note, formats["note"])
    sheet.merge_range(4, 0, 4, max(0, len(contract["fields"]) - 1),
                      "Azul = captura editable. Blanco/null = desconocido; 0 = cero medido. No uses fórmulas.", formats["note"])
    for column, field in enumerate(contract["fields"]):
        sheet.write(5, column, field["name"], formats["header"])
        sheet.set_column(column, column, max(14, min(26, len(field["name"]) + 4)), formats["input"])
        _validation(sheet, column, field)
    for row_index, row in enumerate(rows, start=6):
        for column, field in enumerate(contract["fields"]):
            value = row[field["name"]]
            if value is None:
                sheet.write_blank(row_index, column, None, formats["input"])
            else:
                sheet.write(row_index, column, value, formats["input"])
    key_columns = [next(index for index, field in enumerate(contract["fields"]) if field["name"] == name)
                   for name in contract["primary_key"]]
    criteria = ",".join(
        f"${xlsxwriter.utility.xl_col_to_name(column)}$7:${xlsxwriter.utility.xl_col_to_name(column)}$1006,{xlsxwriter.utility.xl_col_to_name(column)}7"
        for column in key_columns
    )
    first_key = key_columns[0]
    key_letter = xlsxwriter.utility.xl_col_to_name(first_key)
    sheet.data_validation(6, first_key, 5 + MAX_INPUT_ROWS, first_key, {
        "validate": "custom",
        "value": f'=OR({key_letter}7="",COUNTIFS({criteria})=1)',
        "error_title": "Clave duplicada",
        "error_message": "Conserva una sola fila por clave declarada.",
    })
    last_column = xlsxwriter.utility.xl_col_to_name(len(contract["fields"]) - 1)
    sheet.autofilter(f"A6:{last_column}{6 + MAX_INPUT_ROWS}")
    sheet.freeze_panes(6, 0)
    sheet.set_row(0, 28)
    sheet.set_row(1, 25)
    sheet.set_row(2, 24)
    sheet.set_row(3, 24)
    sheet.set_row(4, 24)
    sheet.set_row(5, 32)
    sheet.protect("", {"select_locked_cells": False, "select_unlocked_cells": True,
                         "autofilter": True, "sort": True})
    _setup_page(sheet, marker)


def _build_workbook(path: Path, kind: str) -> None:
    workbook = xlsxwriter.Workbook(path, {"constant_memory": False})
    workbook.set_properties({
        "title": f"Alma OS {CONTRACT_VERSION} {kind}",
        "subject": "Captura local de fuentes operativas agregadas",
        "author": "Alma de Lujo",
        "comments": "Sin macros, sin red, sin datos privados.",
        "created": datetime(2026, 9, 21, tzinfo=timezone.utc),
    })
    workbook.set_calc_mode("auto")
    formats = _formats(workbook)
    _write_support_sheets(workbook, kind, formats)
    rows = synthetic_rows() if kind == "synthetic" else {name: [] for name in SOURCE_NAMES}
    for relation in SOURCE_NAMES:
        _write_relation_sheet(workbook, relation, rows[relation], kind, formats)
    workbook.close()


def _copy_policy(source: Path, destination: Path) -> dict[str, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, destination)
    content = json.loads(destination.read_text(encoding="utf-8"))
    load_policy(destination, as_of="2026-09-21", real_cut=False)
    if content["status"] == "APPROVED" or content["owner_approval_ref"] is not None:
        raise ValueError("public policy must not claim owner approval")
    return {"path": destination.as_posix(), "status": content["status"]}


def build_operating_workbooks(output: str | Path) -> dict[str, Any]:
    """Create a new, public-safe workbook bundle without overwriting output."""

    destination = Path(output)
    if destination.exists() and (not destination.is_dir() or any(destination.iterdir())):
        raise ValueError("output must be a new or empty directory")
    destination.mkdir(parents=True, exist_ok=True)

    workbook_entries: dict[str, dict[str, str]] = {}
    pack_entries: dict[str, dict[str, str]] = {}
    for kind in ("blank", "synthetic"):
        workbook_path = destination / f"operating-v1-{kind}.xlsx"
        pack_path = destination / "source-packs" / kind
        _build_workbook(workbook_path, kind)
        initialize_pack(kind, pack_path)
        if kind == "synthetic":
            parse_pack(pack_path)
        workbook_entries[kind] = {
            "path": workbook_path.relative_to(destination).as_posix(),
            "sha256": _sha256(workbook_path),
        }
        pack_entries[kind] = {"path": pack_path.relative_to(destination).as_posix()}

    policies: list[dict[str, str]] = []
    for filename in POLICY_FILES:
        source = ROOT / "client" / "v1" / "policies" / filename
        target = destination / "policies" / filename
        entry = _copy_policy(source, target)
        entry["path"] = target.relative_to(destination).as_posix()
        entry["sha256"] = _sha256(target)
        policies.append(entry)

    files = {
        path.relative_to(destination).as_posix(): _sha256(path)
        for path in sorted(destination.rglob("*"))
        if path.is_file()
    }
    manifest: dict[str, Any] = {
        "manifest_version": "operating-workbook-manifest-v1",
        "contract_version": CONTRACT_VERSION,
        "relation_count": len(SOURCE_NAMES),
        "relations": list(SOURCE_NAMES),
        "support_sheets": list(SUPPORT_SHEETS),
        "completeness_matrix": [COMPLETENESS_MATRIX[name] for name in SOURCE_NAMES],
        "workbooks": workbook_entries,
        "source_packs": pack_entries,
        "policies": policies,
        "files": files,
        "public_safety": {
            "synthetic_only": True,
            "private_data": False,
            "external_links": False,
            "macros": False,
            "owner_approval_claimed": False,
        },
    }
    (destination / "workbook-manifest.json").write_bytes(canonical_json(manifest) + b"\n")
    return manifest


def build_parity_receipt(workbook_path: str | Path, canonical_pack: str | Path,
                         receipt_path: str | Path) -> dict[str, Any]:
    """Independently prove workbook, adapter pack and manifest parity."""

    from csv import DictReader
    from datetime import date as date_type, datetime as datetime_type

    from openpyxl import load_workbook

    from alma.operating_workbook import export_workbook_to_pack

    workbook_path = Path(workbook_path)
    canonical_pack = Path(canonical_pack)
    workbook = load_workbook(workbook_path, data_only=False, keep_links=False)
    relation_receipts: dict[str, Any] = {}
    with tempfile.TemporaryDirectory(prefix="alma-workbook-oracle-") as tmp:
        private_root = Path(tmp)
        export = export_workbook_to_pack(workbook_path, private_root / "export",
                                         private_root=private_root)
        exported_pack = Path(export["pack_path"])
        export_manifest = json.loads(Path(export["manifest_path"]).read_text(encoding="utf-8"))
        for relation in SOURCE_NAMES:
            fields = [field["name"] for field in SOURCES[relation]["fields"]]
            sheet = workbook[relation]
            raw_rows = [
                [sheet.cell(row, column).value for column in range(1, len(fields) + 1)]
                for row in range(7, sheet.max_row + 1)
            ]
            raw_rows = [row for row in raw_rows if any(value is not None for value in row)]
            with (exported_pack / f"{relation}.csv").open(encoding="utf-8", newline="") as stream:
                exported_rows = list(DictReader(stream))
            expected_rows = [
                ["" if value is None else value.isoformat()
                 if isinstance(value, (date_type, datetime_type)) else str(value) for value in row]
                for row in raw_rows
            ]
            actual_rows = [[row[field] for field in fields] for row in exported_rows]
            exported_hash = _sha256(exported_pack / f"{relation}.csv")
            canonical_hash = _sha256(canonical_pack / f"{relation}.csv")
            units = {field["name"]: field["unit"] for field in SOURCES[relation]["fields"]}
            cents_fields = [name for name, unit in units.items() if unit == "MXN cents"]
            pass_relation = (
                expected_rows == actual_rows
                and exported_hash == canonical_hash
                and exported_hash == export_manifest["relations"][relation]["sha256"]
                and all(value == "" or Decimal(value) == Decimal(value).to_integral_value()
                        for row in exported_rows for name, value in row.items() if name in cents_fields)
            )
            relation_receipts[relation] = {
                "disposition": "PASS" if pass_relation else "BLOCKED",
                "row_count": len(actual_rows),
                "null_count": sum(value == "" for row in actual_rows for value in row),
                "zero_count": sum(value == "0" for row in actual_rows for value in row),
                "cents_fields": cents_fields,
                "units_sha256": hashlib.sha256(canonical_json(units)).hexdigest(),
                "source_sha256": exported_hash,
            }
    workbook.close()
    source_hashes = {name: value["source_sha256"] for name, value in relation_receipts.items()}
    receipt = {
        "receipt_version": "workbook-pack-parity-v1",
        "status": "PASS" if len(relation_receipts) == 22 and
                  all(value["disposition"] == "PASS" for value in relation_receipts.values()) else "BLOCKED",
        "contract_version": CONTRACT_VERSION,
        "synthetic_business_data": True,
        "workbook_sha256": _sha256(workbook_path),
        "pack_sha256": hashlib.sha256(canonical_json(source_hashes)).hexdigest(),
        "manifest_sha256": export["manifest_sha256"],
        "relation_count": len(relation_receipts),
        "relation_ids": list(SOURCE_NAMES),
        "relations": relation_receipts,
        "oracle_command": ".venv\\Scripts\\python.exe -m unittest tests.test_operating_workbooks.WorkbookImportTests -v",
        "assertions": [
            "workbook cells equal exported CSV bytes by relation and ordered field",
            "exported source hashes equal canonical synthetic source-pack hashes",
            "blank/null remains empty, explicit zero remains 0",
            "all MXN cents values are exact integers",
            "manifest units and source hashes cover all 22 relations",
        ],
        "contains_absolute_paths": False,
        "contains_cell_values": False,
    }
    if receipt["status"] != "PASS":
        raise ValueError("workbook parity oracle blocked")
    target = Path(receipt_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(canonical_json(receipt) + b"\n")
    return receipt


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args(argv)
    try:
        manifest = build_operating_workbooks(args.output)
    except (OSError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    print(json.dumps({"status": "PASS", "output": str(args.output),
                      "relations": manifest["relation_count"]}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
