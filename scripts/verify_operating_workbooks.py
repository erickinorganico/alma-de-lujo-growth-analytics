#!/usr/bin/env python3
"""Independent acceptance boundary for the two operating-v1 delivery books.

The books are capture surfaces. Their only formulas are data-validation rules;
there are no calculated cells, cached business results, or displayed totals.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import posixpath
import re
import shutil
import sys
import tempfile
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from alma.operating_contracts import CONTRACT_VERSION, SOURCE_NAMES, SOURCES, canonical_json, validate_metadata  # noqa: E402
from alma.operating_interchange import parse_pack  # noqa: E402
from alma.operating_workbook import export_workbook_to_pack  # noqa: E402

SUPPORT = ("INICIO", "DICCIONARIO", "COMPLETITUD")
FILES = {
    "blank": "Alma_de_Lujo_OPERACION_PLANTILLA.xlsx",
    "synthetic": "Alma_de_Lujo_OPERACION_EJEMPLO.xlsx",
}
INTEGER_TYPES = {"integer", "nonnegative_integer", "signed_integer"}
LEAK = re.compile(r"(?:[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|(?:[A-Za-z]:\\|/Users/|/home/|file://)|\b\+?\d[\d ()-]{8,}\d\b)", re.I)
CELL_LEAK = re.compile(r"(?:[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}|(?:[A-Za-z]:\\|/Users/|/home/|file://))", re.I)
ACTIVE_PART = re.compile(r"(?:vbaProject|externalLinks|embeddings|connections|customXml|activeX|oleObjects|queryTables|pivotCache|powerPivot)", re.I)
SAFE_PART = re.compile(r"(?:\[Content_Types\]\.xml|_rels/\.rels|xl/_rels/workbook\.xml\.rels|xl/workbook\.xml|xl/worksheets/sheet\d+\.xml|xl/sharedStrings\.xml|xl/styles\.xml|xl/theme/theme\d+\.xml|docProps/(?:core|app)\.xml)")
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


class WorkbookAcceptanceError(ValueError):
    pass


def require(condition: bool, message: str) -> None:
    if not condition:
        raise WorkbookAcceptanceError(message)


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha(path: Path) -> str:
    return digest(path.read_bytes())


def normalized(value: Any, field: dict[str, Any], where: str) -> str:
    if value is None or value == "":
        return ""
    require(not isinstance(value, bool), f"{where}: boolean value")
    kind = field["type"]
    if kind in INTEGER_TYPES:
        try:
            number = Decimal(str(value))
        except (InvalidOperation, ValueError) as exc:
            raise WorkbookAcceptanceError(f"{where}: invalid integer") from exc
        require(number.is_finite() and number == number.to_integral_value(), f"{where}: fractional/nonfinite integer")
        if kind == "nonnegative_integer":
            require(number >= 0, f"{where}: negative integer")
        # Decimal is deliberate: binary float and rounding are never money authority.
        return format(number.quantize(Decimal(1)), "f")
    if kind == "date" and isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if kind == "timestamp" and isinstance(value, datetime):
        require(value.tzinfo is not None, f"{where}: timestamp without timezone")
        return value.isoformat()
    require(isinstance(value, str), f"{where}: text required")
    require(not value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")), f"{where}: formula injection")
    return value


def csv_bytes(relation: str, rows: list[list[str]]) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, lineterminator="\n")
    writer.writerow([field["name"] for field in SOURCES[relation]["fields"]])
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def archive_inventory(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        require(len(names) == len(set(names)), "duplicate ZIP member")
        require(len(names) < 400 and sum(item.file_size for item in archive.infolist()) < 50_000_000,
                "oversized XLSX package")
        result: dict[str, str] = {}
        targets: list[tuple[str, str]] = []
        for item in archive.infolist():
            name = item.filename.replace("\\", "/")
            require(name == item.filename and not name.startswith("/") and ".." not in name.split("/"),
                    f"unsafe ZIP member: {name}")
            require(not (item.flag_bits & 1) and not (item.external_attr >> 16) & 0o170000 == 0o120000,
                    f"encrypted/symlink ZIP member: {name}")
            require(not ACTIVE_PART.search(name), f"active XLSX part: {name}")
            require(bool(SAFE_PART.fullmatch(name)), f"unexpected XLSX part: {name}")
            data = archive.read(item)
            if name.endswith(".rels"):
                root = ET.fromstring(data)
                ids: set[str] = set()
                for relation in root.findall(f"{REL_NS}Relationship"):
                    relation_id = relation.attrib.get("Id", "")
                    require(relation_id and relation_id not in ids, f"duplicate relationship ID: {name}")
                    ids.add(relation_id)
                    require(relation.attrib.get("TargetMode") != "External", f"external relationship: {name}")
                    target = relation.attrib.get("Target", "")
                    require(not ("://" in target or target.startswith("file:")), f"external relationship target: {name}")
                    source = name.replace("/_rels/", "/").removesuffix(".rels")
                    if name == "_rels/.rels":
                        source = ""
                    resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), target.lstrip("/")))
                    if target.startswith("/"):
                        resolved = target.lstrip("/")
                    targets.append((name, resolved))
            if name == "[Content_Types].xml":
                require(b"macroEnabled" not in data and b"vbaProject" not in data, "macro content type")
            if name.startswith("docProps/"):
                require(not LEAK.search(data.decode("utf-8", errors="replace")), f"metadata leak: {name}")
            result[name] = digest(data)
        require("xl/workbook.xml" in result and "[Content_Types].xml" in result,
                "missing core XLSX parts")
        for origin, target in targets:
            require(target in result, f"missing relationship target: {origin} -> {target}")
        return dict(sorted(result.items()))


def validation_inventory(sheet: Any) -> list[dict[str, Any]]:
    return sorted(({
        "range": str(rule.sqref),
        "type": rule.type,
        "operator": rule.operator,
        "formula1": rule.formula1,
        "formula2": rule.formula2,
        "allow_blank": rule.allow_blank,
        "error": rule.error,
    } for rule in sheet.data_validations.dataValidation), key=lambda row: (row["range"], row["type"] or ""))


def inspect_book(path: Path, kind: str) -> dict[str, Any]:
    members = archive_inventory(path)
    with tempfile.TemporaryDirectory(prefix="alma-workbook-accept-") as scratch:
        private = Path(scratch)
        parsed = None
        book = load_workbook(path, data_only=False, keep_links=False)
        cached = load_workbook(path, data_only=True, keep_links=False)
        try:
            require(book.sheetnames == [*SUPPORT, *SOURCE_NAMES], "sheet order or relation set")
            require(book["INICIO"]["B4"].value == ("BLANK" if kind == "blank" else "SYNTHETIC_EXAMPLE"),
                    "workbook kind")
            require(book["INICIO"]["B3"].value == CONTRACT_VERSION, "contract version")
            def optional(value: Any) -> Any:
                return None if value == "POR_COMPLETAR" else value
            metadata = {
                "contract_version": book["INICIO"]["B3"].value,
                "input_class": book["INICIO"]["B4"].value,
                "cutoff_at": optional(book["INICIO"]["B5"].value),
                "timezone": optional(book["INICIO"]["B6"].value),
                "currency": book["INICIO"]["B7"].value,
                "coverage": {
                    relation: {
                        "status": book["COMPLETITUD"].cell(index, 6).value,
                        "window_start": book["COMPLETITUD"].cell(index, 7).value,
                        "window_end": book["COMPLETITUD"].cell(index, 8).value,
                    }
                    for index, relation in enumerate(SOURCE_NAMES, start=2)
                },
            }
            require([book["COMPLETITUD"].cell(index, 1).value for index in range(2, 24)] == list(SOURCE_NAMES),
                    "completeness relation order")
            validate_metadata(metadata, allow_blank=kind == "blank")
            require(not book._external_links, "external workbook links")
            require(not list(book.defined_names.values()) or all(name.name.startswith("_xlnm.") for name in book.defined_names.values()),
                    "custom defined name")
            props = book.properties
            require(props.creator == "Alma de Lujo" and props.lastModifiedBy in (None, "Alma de Lujo"),
                    "unexpected workbook author metadata")
            require(props.title == f"Alma OS {CONTRACT_VERSION} {kind}", "unexpected workbook title metadata")
            for value in (props.creator, props.lastModifiedBy, props.keywords, props.description, props.identifier, props.category):
                require(not (isinstance(value, str) and LEAK.search(value)), "metadata leak")
            sheets: dict[str, Any] = {}
            formula_cells: dict[str, dict[str, str]] = {}
            source_hashes: dict[str, str] = {}
            row_counts: dict[str, int] = {}
            type_inventory: dict[str, Any] = {}
            expected_csv_files: dict[str, bytes] = {}
            for name in book.sheetnames:
                sheet = book[name]
                require(sheet.sheet_state == "visible", f"hidden sheet: {name}")
                require(not any(dim.hidden for dim in sheet.row_dimensions.values()), f"hidden row: {name}")
                require(not any(dim.hidden for dim in sheet.column_dimensions.values()), f"hidden column: {name}")
                formulas: dict[str, str] = {}
                cells: list[list[Any]] = []
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            cells.append([cell.coordinate, cell.data_type, str(cell.value)])
                        if isinstance(cell.value, str) and cell.data_type != "f":
                            require(not CELL_LEAK.search(cell.value), f"cell privacy leak: {name}!{cell.coordinate}")
                            require(not cell.value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")),
                                    f"formula injection: {name}!{cell.coordinate}")
                        require(cell.comment is None, f"comment: {name}!{cell.coordinate}")
                        require(cell.data_type != "e", f"error cell: {name}!{cell.coordinate}")
                        if cell.data_type == "f":
                            formulas[cell.coordinate] = str(cell.value)
                        if cell.hyperlink is not None:
                            location = cell.hyperlink.location or ""
                            require(cell.hyperlink.target is None and any(location == f"'{relation}'!A1" for relation in SOURCE_NAMES),
                                    f"external hyperlink: {name}!{cell.coordinate}")
                require(not formulas, f"unexpected formula cell: {name}")
                # With no formulas, cached values must equal literal values at every cell.
                for row in sheet.iter_rows():
                    for cell in row:
                        saved = cached[name][cell.coordinate]
                        require(saved.value == cell.value, f"altered cached result: {name}!{cell.coordinate}")
                formula_cells[name] = formulas
                sheets[name] = {
                    "dimension": sheet.calculate_dimension(),
                    "max_row": sheet.max_row,
                    "max_column": sheet.max_column,
                    "cell_sha256": digest(canonical_json(cells)),
                    "tables": sorted(sheet.tables),
                    "validations": validation_inventory(sheet),
                    "print_area": str(sheet.print_area) if sheet.print_area else None,
                    "print_title_rows": sheet.print_title_rows,
                    "auto_filter": sheet.auto_filter.ref,
                }
                if name not in SOURCE_NAMES:
                    continue
                fields = SOURCES[name]["fields"]
                require([sheet.cell(6, index).value for index in range(1, len(fields) + 1)] ==
                        [field["name"] for field in fields], f"source header: {name}")
                require(sheet.max_column == len(fields), f"source width: {name}")
                require(len(sheet.data_validations.dataValidation) >= len(fields), f"missing validations: {name}")
                values: list[list[str]] = []
                for row_num in range(7, sheet.max_row + 1):
                    raw = [sheet.cell(row_num, col).value for col in range(1, len(fields) + 1)]
                    if all(value is None or value == "" for value in raw):
                        continue
                    values.append([normalized(value, field, f"{name}!{row_num}:{field['name']}")
                                   for value, field in zip(raw, fields, strict=True)])
                expected_csv = csv_bytes(name, values)
                expected_csv_files[name] = expected_csv
                source_hashes[name] = digest(expected_csv)
                row_counts[name] = len(values)
                type_inventory[name] = {
                    field["name"]: {"type": field["type"], "unit": field["unit"],
                                    "null_count": sum(row[index] == "" for row in values),
                                    "zero_count": sum(row[index] == "0" for row in values)}
                    for index, field in enumerate(fields)
                }
            require(kind != "blank" or all(count == 0 for count in row_counts.values()), "blank book contains source rows")
            require(kind != "synthetic" or all(count > 0 for count in row_counts.values()), "example lacks relation rows")
            if kind == "synthetic":
                exported = export_workbook_to_pack(path, private / "export", private_root=private)
                pack = Path(exported["pack_path"])
                parsed = parse_pack(pack)
                for name in SOURCE_NAMES:
                    require(expected_csv_files[name] == (pack / f"{name}.csv").read_bytes(),
                            f"independent CSV/adapter mismatch: {name}")
                    require(len(parsed["tables"][name]) == row_counts[name], f"intake row mismatch: {name}")
            return {
                "path": path.name,
                "sha256": sha(path),
                "members": members,
                "sheets": sheets,
                "formula_cells": formula_cells,
                "formula_count": 0,
                "cached_result_count": 0,
                "business_total_count": 0,
                "source_pack_sha256": source_hashes,
                "row_counts": row_counts,
                "field_types_and_counts": type_inventory,
                "metadata_sha256": digest(canonical_json(metadata) + b"\n"),
                "normalized_rows_digest": parsed["normalized_rows_digest"] if parsed else None,
            }
        finally:
            cached.close()
            book.close()


def build_manifest(directory: Path) -> dict[str, Any]:
    return {
        "manifest_version": "operating-workbook-acceptance-v1",
        "contract_version": CONTRACT_VERSION,
        "stage": "PRE_EXCEL",
        "relations": list(SOURCE_NAMES),
        "support_sheets": list(SUPPORT),
        "calculated_business_results": False,
        "workbooks": {kind: inspect_book(directory / filename, kind) for kind, filename in FILES.items()},
    }


def check_manifest(path: Path) -> dict[str, Any]:
    expected = json.loads(path.read_text(encoding="utf-8"))
    require(expected["manifest_version"] == "operating-workbook-acceptance-v1", "manifest version")
    require(expected["relations"] == list(SOURCE_NAMES), "manifest relations")
    require(expected["stage"] in ("PRE_EXCEL", "FINAL_EXCEL"), "manifest stage")
    actual = build_manifest(path.parent)
    actual["stage"] = expected["stage"]
    require(actual == expected, "workbook manifest/hash/surface mismatch")
    return actual


def pdf_page_count(path: Path) -> int:
    data = path.read_bytes()
    require(data.startswith(b"%PDF-"), f"invalid PDF: {path.name}")
    count = len(re.findall(rb"/Type\s*/Page\b", data))
    require(count > 0, f"PDF has no countable pages: {path.name}")
    return count


def check_receipts(manifest: dict[str, Any], excel_path: Path | None, visual_path: Path | None) -> dict[str, int]:
    require(manifest["stage"] == "FINAL_EXCEL", "final Excel stage required")
    require(excel_path is not None and visual_path is not None, "both Excel and visual receipts required")
    require(excel_path.is_file() and visual_path.is_file(), "Excel or visual receipt missing")
    root = excel_path.parent.resolve()
    require(visual_path.parent.resolve() == root, "receipts must share evidence directory")
    excel = json.loads(excel_path.read_text(encoding="utf-8"))
    visual = json.loads(visual_path.read_text(encoding="utf-8"))
    require(excel.get("schema") == "operating-excel-recalculation-v1", "Excel receipt schema")
    require(visual.get("schema") == "operating-visual-inspection-v1", "visual receipt schema")
    for receipt in (excel, visual):
        require(isinstance(receipt.get("timestamp_utc"), str) and
                bool(re.fullmatch(r"\d{4}-\d\d-\d\dT\d\d:\d\d:\d\dZ", receipt["timestamp_utc"])),
                "receipt UTC timestamp")
    require(visual.get("excel_receipt_sha256") == sha(excel_path), "stale visual/Excel receipt")
    require(visual.get("inspector") and visual.get("runtime"), "visual inspector/runtime missing")
    expected_files = {excel_path.name, visual_path.name}
    expected_pages: dict[tuple[str, str, int], dict[str, Any]] = {}
    books = excel.get("workbooks")
    require(isinstance(books, list) and len(books) == len(FILES), "Excel book count")
    require([entry.get("kind") for entry in books] == list(FILES), "Excel book order/kinds")
    for book in books:
        kind = book["kind"]
        expected = manifest["workbooks"][kind]
        require(book.get("file") == FILES[kind] and book.get("sha256") == expected["sha256"],
                f"stale Excel workbook: {kind}")
        require(book.get("engine") == "Microsoft Excel" and book.get("version") and book.get("build"),
                f"Excel engine/version/build: {kind}")
        require(book.get("calculate_full_rebuild") is True and book.get("calculation_state") == 0 and
                book.get("calculation_mode") in (-4105, -4135, 2), f"Excel calculation state/mode: {kind}")
        sheets = book.get("sheets")
        printed = {name: row["print_area"].split("!", 1)[-1]
                   for name, row in expected["sheets"].items() if row["print_area"]}
        require(isinstance(sheets, list) and len(sheets) == len(printed), f"printed sheet count: {kind}")
        names = [entry.get("sheet") for entry in sheets]
        require(len(names) == len(set(names)) and set(names) == set(printed),
                f"printed sheet set: {kind}")
        for sheet in sheets:
            name = sheet["sheet"]
            filename = f"{kind}__{name}.pdf"
            require(sheet.get("print_area") == printed[name] and sheet.get("pdf") == filename,
                    f"print area/PDF identity: {kind}/{name}")
            pdf = root / filename
            require(pdf.is_file() and sheet.get("pdf_sha256") == sha(pdf), f"PDF missing/hash: {filename}")
            expected_files.add(filename)
            count = pdf_page_count(pdf)
            for number in range(1, count + 1):
                key = (kind, filename, number)
                png = f"rendered-pages/{kind}__{name}__page-{number:02d}.png"
                expected_pages[key] = {
                    "kind": kind, "workbook_sha256": expected["sha256"], "pdf": filename,
                    "pdf_sha256": sheet["pdf_sha256"], "page": number, "page_count": count,
                    "png": png,
                }
                expected_files.add(png)
    actual_pages = visual.get("pages")
    require(isinstance(actual_pages, list) and len(actual_pages) == len(expected_pages), "visual page count")
    seen: set[tuple[str, str, int]] = set()
    for page in actual_pages:
        require(isinstance(page, dict), "invalid visual page")
        key = (page.get("kind"), page.get("pdf"), page.get("page"))
        require(key in expected_pages and key not in seen, "missing/duplicate/unexpected visual page")
        seen.add(key)
        expected = expected_pages[key]
        require(all(page.get(field) == value for field, value in expected.items()), "stale visual page binding")
        png = root / expected["png"]
        require(png.is_file() and png.read_bytes().startswith(b"\x89PNG\r\n\x1a\n") and
                page.get("png_sha256") == sha(png), f"PNG missing/hash: {expected['png']}")
        require(page.get("inspected") is True and page.get("disposition") == "PASS" and
                page.get("legibility") == "READABLE" and page.get("clipping") == "NONE",
                f"uninspected/illegible/clipped page: {expected['png']}")
    require(seen == set(expected_pages), "visual page inventory incomplete")
    actual_files = {p.relative_to(root).as_posix() for p in root.rglob("*") if p.is_file()}
    require(actual_files == expected_files, f"unexpected/missing Excel evidence files: {sorted(actual_files ^ expected_files)}")
    return {"workbooks": len(books), "pdfs": len(expected_files) - len(expected_pages) - 2,
            "pages": len(expected_pages)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("build", "check"))
    parser.add_argument("--manifest", type=Path, default=ROOT / "client/v1/workbook-manifest.json")
    parser.add_argument("--excel-receipt", type=Path)
    parser.add_argument("--visual-receipt", type=Path)
    args = parser.parse_args()
    try:
        if args.command == "build":
            args.manifest.write_bytes(canonical_json(build_manifest(args.manifest.parent)) + b"\n")
        else:
            manifest = check_manifest(args.manifest)
            if manifest["stage"] == "FINAL_EXCEL":
                check_receipts(manifest, args.excel_receipt, args.visual_receipt)
        print(json.dumps({"status": "PASS", "stage": args.command, "manifest": str(args.manifest)}, sort_keys=True))
        return 0
    except (OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        print(f"BLOCKED: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
